#!/usr/bin/env python3
"""
NOAA GSL MATS Automated Verification & Presentation Generator (mats_reporter)
Author: RRFS Development Pair
"""

import argparse
import asyncio
import base64
import datetime
import html
import os
import re
import shutil
import subprocess
import sys
import yaml
import json
import shlex
import tempfile
import functools
from PIL import Image
from playwright.async_api import async_playwright

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(line_buffering=True)
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(line_buffering=True)

# Ensure all stdout prints are unbuffered and flushed immediately across pipes
print = functools.partial(print, flush=True)

# Prevent PYTHONHOME/PYTHONPATH from poisoning any child Python processes (e.g. hpc-run)
os.environ.pop("PYTHONHOME", None)
os.environ.pop("PYTHONPATH", None)

VALID_HPC_CLUSTERS = {"gaea", "gaea-c5", "gaea-c6", "gaea-c7", "hera", "ursa", "orion", "hercules"}

def reject_pty_control_chars(s: str) -> str:
    """
    Guards against PTY line discipline injection attacks by rejecting
    ASCII C0 control characters (0x00-0x1F, including newline/carriage return) and 0x7F (DEL).
    Allows all valid path characters (spaces, brackets, unicode, etc.) to be safely handled by shlex.
    """
    if not s:
        return ""
    if any(ord(ch) < 32 or ord(ch) == 127 for ch in s):
        raise ValueError(f"Security Alert: Input contains illegal PTY control character(s): {repr(s)}")
    return s

def report_progress(total_pct: int, stage_pct: int = None, stage_title: str = ""):
    """
    Outputs a structured dual progress marker consumed by Verification Studio GUI.
    Format: [PROGRESS: <total_pct>%, stage=<stage_pct>%] [<stage_title>]
    """
    t_pct = max(0, min(100, int(total_pct)))
    s_pct = max(0, min(100, int(stage_pct))) if stage_pct is not None else t_pct
    print(f"[PROGRESS: {t_pct}%, stage={s_pct}%] [{stage_title}]", flush=True)

class PipelineProgressTracker:
    """
    Dynamic Stage Planner & Dual-Tier Progress Coordinator.
    Computes smooth cumulative total progress (0% ~ 100%) and stage progress (0% ~ 100%).
    """
    def __init__(self, stages):
        # stages is a list of (stage_id, stage_name, weight)
        self.stages = stages
        self.total_stages = len(stages)
        total_weight = sum(w for _, _, w in stages) or 1.0
        self.bounds = []
        accum = 0.0
        for _, _, w in stages:
            start = (accum / total_weight) * 100.0
            accum += w
            end = (accum / total_weight) * 100.0
            self.bounds.append((start, end))

    def update(self, stage_idx: int, stage_pct: int):
        stage_pct = max(0, min(100, int(stage_pct)))
        stage_num = stage_idx + 1
        _, stage_name, _ = self.stages[stage_idx]
        stage_title = f"Stage {stage_num}/{self.total_stages}: {stage_name}"
        start_pct, end_pct = self.bounds[stage_idx]
        total_pct = int(start_pct + (end_pct - start_pct) * (stage_pct / 100.0))
        report_progress(total_pct, stage_pct, stage_title)

    def complete(self):
        report_progress(100, 100, "Complete: Verification Report Ready")


def is_valid_image(path: str, min_size: int = 4096) -> bool:
    """
    Verifies image integrity using binary magic bytes and Pillow structural decoding.
    Rejects truncated files, zero-byte stubs, and HTML error responses (e.g. 502/504 Bad Gateway).
    """
    if not path or not os.path.exists(path):
        return False
    try:
        if os.path.getsize(path) < min_size:
            return False
        with open(path, "rb") as f:
            header = f.read(8)
            # PNG magic: \x89PNG\r\n\x1a\n ; JPEG magic: \xff\xd8\xff
            if not (header.startswith(b"\x89PNG\r\n\x1a\n") or header.startswith(b"\xff\xd8\xff")):
                return False
        with Image.open(path) as img:
            img.verify()
        return True
    except Exception:
        return False

def atomic_write_bytes(target_path: str, data: bytes) -> None:
    """
    Atomically writes binary data to target_path using POSIX rename(2).
    Eliminates TOCTOU race conditions and never follows symlinks on destination.
    """
    target_dir = os.path.dirname(os.path.abspath(target_path))
    os.makedirs(target_dir, exist_ok=True)
    tmp_fd, tmp_path = tempfile.mkstemp(dir=target_dir, prefix=".tmp_write_", suffix=".bin")
    try:
        with os.fdopen(tmp_fd, "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp_path, target_path)
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass
        raise

def atomic_write_text(target_path: str, text: str, encoding: str = "utf-8") -> None:
    """
    Atomically writes text to target_path.
    """
    atomic_write_bytes(target_path, text.encode(encoding))

def atomic_save_pptx(prs, output_pptx: str) -> None:
    """
    Atomically saves python-pptx presentation to output_pptx.
    """
    target_dir = os.path.dirname(os.path.abspath(output_pptx))
    os.makedirs(target_dir, exist_ok=True)
    tmp_fd, tmp_path = tempfile.mkstemp(dir=target_dir, prefix=".tmp_pptx_", suffix=".pptx")
    os.close(tmp_fd)
    try:
        prs.save(tmp_path)
        os.replace(tmp_path, output_pptx)
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass
        raise

DEFAULT_PALETTE = [
    "rgb(255, 0, 0)",      # Red
    "rgb(0, 0, 255)",      # Blue
    "rgb(255, 165, 0)",    # Orange
    "rgb(128, 128, 128)",  # Grey
    "rgb(46, 139, 87)",    # Sea Green
    "rgb(148, 0, 211)"     # Dark Violet
]

REGION_MAP = {
    "surface": {
        "all": "All HRRR domain",
        "east": "Eastern HRRR domain",
        "west": "Western HRRR domain"
    },
    "ceil_vis": {
        "all": "All HRRR domain",
        "east": "Eastern HRRR domain",
        "west": "Western HRRR domain"
    },
    "upperair": {
        "all": "HRRR domain",
        "east": "East of 109W",
        "west": "West of 109W"
    },
    "radar": {
        "all": "Continental US",
        "east": "Eastern US (lon <= 100W)",
        "west": "Western US (lon >= 100W)"
    }
}

REGION_NAMES = {
    "all": "Full Domain (CONUS)",
    "east": "Eastern Region (East of 109°W / lon ≤ 100°W)",
    "west": "Western Region (Complex Terrain / West of 109°W)"
}

def extract_short_name(model_id):
    """
    Intelligently extracts concise short name from MATS model ID.
    e.g. 'RRFSv2_conus_3km_ret_generictquv' -> 'generictquv'
         'RRFSv2_conus_3km_ret_halfTerr'    -> 'halfTerr'
         'RRFSv2_conus_3km_ret_RRFSv2X_retv3' -> 'retv3'
         'RRFSv2_conus_3km_ret_v2x_retv4'   -> 'retv4'
    """
    mid = model_id.strip()
    if '_ret_' in mid:
        name = mid.split('_ret_', 1)[1]
    elif '_ret' in mid:
        name = mid.split('_ret', 1)[1].lstrip('_')
    else:
        name = mid.split('_')[-1]
    if 'retv' in name.lower():
        idx = name.lower().find('retv')
        name = name[idx:]
    name = re.sub(r'[^a-zA-Z0-9._-]', '_', name)
    return name

def normalize_hpc_cluster(cluster, path):
    """
    Intelligently routes and normalizes machine names to their canonical HPC cluster ID.
    Supports: gaeac5, gaeac6, gaeac7, gaea-c5, gaea-c6, gaea-c7, c5, c6, c7, gaea, hera, ursa, orion, hercules.
    Auto-detects /gpfs/f6 -> gaea-c6, /gpfs/f7 -> gaea-c7, /gpfs/f5 -> gaea-c5.
    """
    alias_map = {
        "gaeac5": "gaea-c5",
        "gaeac6": "gaea-c6",
        "gaeac7": "gaea-c7",
        "gaea-c5": "gaea-c5",
        "gaea-c6": "gaea-c6",
        "gaea-c7": "gaea-c7",
        "c5": "gaea-c5",
        "c6": "gaea-c6",
        "c7": "gaea-c7",
    }
    if not cluster:
        if path:
            if "/gpfs/f6" in path:
                return "gaea-c6"
            elif "/gpfs/f7" in path:
                return "gaea-c7"
            elif "/gpfs/f5" in path:
                return "gaea-c5"
            elif "/scratch1/RDARCH" in path or "/scratch2" in path:
                return "hera"
            elif "/work/noaa" in path or "/work2/noaa" in path:
                return "orion"
        return cluster

    c = cluster.lower().strip()
    if c in alias_map:
        return alias_map[c]

    if c == "gaea":
        if path:
            if "/gpfs/f6" in path:
                return "gaea-c6"
            elif "/gpfs/f7" in path:
                return "gaea-c7"
            elif "/gpfs/f5" in path:
                return "gaea-c5"
        return "gaea-c5"
    return c


def parse_model_entry(entry):
    """
    Parses flexible model input entry:
    Formats supported:
    1. 'ID - cluster:path' (e.g. 'RRFSv2_conus_3km_ret_generictquv - gaea:/gpfs/f6/...')
    2. 'ID - path'         (local or default cluster path)
    3. 'ID'                (MATS ID only, e.g. 'RRFSv2_conus_3km_ret_v2x_retv4')
    4. 'name : ID - ...'   (backwards compatible with custom name)
    """
    if isinstance(entry, dict):
        mid = entry.get('id') or entry.get('model_id') or ''
        name = entry.get('name') or extract_short_name(mid)
        path = entry.get('path')
        cluster = normalize_hpc_cluster(entry.get('cluster'), path)
        user = entry.get('user')
        color = entry.get('color')
        return {'id': mid, 'name': name, 'cluster': cluster, 'path': path, 'user': user, 'color': color}

    entry = str(entry).strip()
    custom_name = None
    if ':' in entry and ' - ' in entry:
        parts_dash = entry.split(' - ', 1)
        if ':' in parts_dash[0] and not ('/' in parts_dash[0].split(':', 1)[0]):
            p_name = parts_dash[0].split(':', 1)
            custom_name = p_name[0].strip()
            entry = p_name[1].strip() + ' - ' + parts_dash[1].strip()
    elif ':' in entry and not (' - ' in entry) and not ('/' in entry.split(':', 1)[0]):
        p_name = entry.split(':', 1)
        custom_name = p_name[0].strip()
        entry = p_name[1].strip()

    cluster = None
    path = None
    user = None
    if ' - ' in entry:
        model_id, path_spec = entry.split(' - ', 1)
        model_id = model_id.strip()
        path_spec = path_spec.strip()
        if ':' in path_spec and not path_spec.startswith('/'):
            c_part, p_part = path_spec.split(':', 1)
            c_part = c_part.strip()
            if '@' in c_part:
                user_part, c_part = c_part.split('@', 1)
                user = user_part.strip()
            cluster = c_part.strip().lower()
            path = p_part.strip()
        else:
            path = path_spec
    else:
        model_id = entry

    cluster = normalize_hpc_cluster(cluster, path)
    name = custom_name if custom_name else extract_short_name(model_id)
    name = re.sub(r'[^a-zA-Z0-9._-]', '_', name)
    return {'id': model_id, 'name': name, 'cluster': cluster, 'path': path, 'user': user}


DEFAULT_RADAR_CASES = []


def get_radar_cases(config):
    """
    Resolves the list of radar comparison cases from configuration.
    Supports configured radar_cases (list of {date, cycle, fhrs}).
    If none configured, returns empty list.
    """
    if config.get("skip_cref"):
        return []
    configured = config.get("radar_cases")
    if configured:
        out_cases = []
        idx = 1
        for item in configured:
            if isinstance(item, (list, tuple)) and len(item) == 5:
                out_cases.append(item)
                continue
            cdate = str(item.get("date", "")).replace("-", "")
            cycle = str(item.get("cycle", "")).zfill(2)
            fhrs = item.get("fhrs", [])
            try:
                dt_obj = datetime.datetime.strptime(cdate, "%Y%m%d")
                mon_day = f"{dt_obj.strftime('%b')} {dt_obj.day}"
            except Exception:
                mon_day = cdate
            for f in fhrs:
                fhr = int(f)
                case_title = f"{fhr}-h fcst from {cycle}Z {mon_day}"
                out_cases.append((str(idx), cdate, cycle, fhr, case_title))
                idx += 1
        return out_cases
    return []



def get_hpc_run_executable():
    """
    Locates the core hpc-run executable.
    Checks environment variable, bundled application framework path,
    system PATH, and user local bin.
    """
    env_path = os.environ.get("HPC_RUN_BIN")
    if env_path and os.path.exists(env_path):
        return env_path
    script_dir = os.path.dirname(os.path.abspath(__file__))
    repo_local = os.path.abspath(os.path.join(script_dir, "hpc-tools", "hpc-run"))
    if os.path.exists(repo_local) and os.access(repo_local, os.X_OK):
        return repo_local
    bundled = os.path.abspath(os.path.join(script_dir, "..", "Frameworks", "python", "bin", "hpc-run"))
    if os.path.exists(bundled):
        return bundled
    found = shutil.which("hpc-run")
    if found:
        return found
    user_local = os.path.expanduser("~/.local/bin/hpc-run")
    if os.path.exists(user_local):
        return user_local
    return "hpc-run"


def is_hpc_online(cluster):
    """
    Checks if target cluster daemon is currently online via hpc-run --status.
    """
    hpc_bin = get_hpc_run_executable()
    target = normalize_hpc_cluster(cluster, None)
    if target == "gaea":
        target = "gaea-c5"
    try:
        res = subprocess.run([hpc_bin, "--status", target], capture_output=True, text=True, timeout=5)
        if res.returncode == 0 and "ONLINE" in res.stdout:
            return True
    except Exception:
        pass
    return False


def find_suggested_hpc_user(target, models=None):
    """
    Finds a recommended HPC username from models, environment, or saved configurations.
    """
    cluster_env = target.upper().replace('-', '_')
    env_u = os.environ.get(f"HPC_{cluster_env}_USER") or os.environ.get("HPC_USER")
    if env_u and env_u.strip():
        return env_u.strip()

    for uf in [os.path.expanduser(f"~/.hpc_user_{target}"), os.path.expanduser("~/.hpc_user")]:
        if os.path.exists(uf):
            try:
                with open(uf, "r", encoding="utf-8") as f:
                    u = f.read().strip()
                    if u:
                        return u
            except Exception:
                pass

    if models:
        for m in models:
            if m.get("user"):
                return m.get("user")
            path = m.get("path") or ""
            m_cluster = m.get("cluster") or ""
            if target in m_cluster or (target.startswith("gaea") and m_cluster.startswith("gaea")):
                m_match = re.search(r'/scratch/([^/]+)/', path)
                if m_match:
                    return m_match.group(1)

    ssh_cfg = os.path.expanduser("~/.ssh/config")
    if os.path.exists(ssh_cfg):
        try:
            with open(ssh_cfg, "r", encoding="utf-8", errors="ignore") as f:
                matched = False
                for line in f:
                    line = line.strip()
                    if line.lower().startswith("host ") and not line.startswith("#"):
                        hosts = [h.strip().lower() for h in line[5:].split() if h.strip() and not h.startswith("*")]
                        matched = any(target in h or h in target for h in hosts)
                    elif matched and line.lower().startswith("user ") and not line.startswith("#"):
                        return line[5:].strip()
        except Exception:
            pass

    return ""


def close_terminal_by_tty(tty_file):
    """
    Closes the macOS Terminal window corresponding to the recorded TTY.
    """
    if not os.path.exists(tty_file):
        return
    try:
        with open(tty_file, "r") as f:
            target_tty = f.read().strip()
        if not target_tty:
            return
        apple_script = f'''
        tell application "Terminal"
            repeat with w in windows
                repeat with t in tabs of w
                    if tty of t is "{target_tty}" then
                        if (count of tabs of w) > 1 then
                            close t saving no
                        else
                            close w saving no
                        end if
                        exit repeat
                    end if
                end repeat
            end repeat
        end tell
        '''
        subprocess.run(["osascript", "-e", apple_script], capture_output=True, timeout=4)
    except Exception:
        pass


def ensure_hpc_connection(cluster, models=None):
    """
    Ensures connection to the target HPC cluster.
    If offline, pops up an interactive macOS Terminal window prompting for username (if needed)
    and running 'hpc-run --start <cluster>' for user PIN + YubiKey authentication.
    Automatically closes the Terminal window upon successful connection using the exact TTY handle.
    """
    target = normalize_hpc_cluster(cluster, None)
    if target == "gaea":
        target = "gaea-c5"

    if target not in VALID_HPC_CLUSTERS:
        print(f"[ERROR] Unsupported or invalid HPC cluster target: {target}", flush=True)
        return False

    if is_hpc_online(target):
        return True

    hpc_bin = get_hpc_run_executable()
    suggested_user = find_suggested_hpc_user(target, models) or ""
    # Reject newlines, NUL bytes, spaces, or shell control characters in username
    if not re.match(r'^[a-zA-Z0-9._-]+$', suggested_user):
        suggested_user = ""

    print(f"\n[HPC] Cluster '{target}' is offline or requires authentication.", flush=True)
    print(f"[HPC] Launching interactive Terminal for PIN + YubiKey authentication...", flush=True)
    print(f"[HPC] Please enter your credentials in the popup Terminal window (or close it to cancel).", flush=True)

    import time
    auth_tempdir = tempfile.mkdtemp(prefix=f"verif_auth_{target}_")
    try:
        os.chmod(auth_tempdir, 0o700)
    except Exception:
        pass

    status_file = os.path.join(auth_tempdir, "auth.status")
    cmd_file = os.path.join(auth_tempdir, "auth.command")
    tty_file = os.path.join(auth_tempdir, "auth.tty")

    window_title = f"MATS_HPC_AUTH_{target.upper()}_{os.getpid()}"
    target_env_suffix = target.upper().replace('-', '_')

    # Safe quoting without outer double quotes to prevent syntax escape
    script_content = f"""#!/bin/bash
unset PYTHONHOME
unset PYTHONPATH
CURRENT_TTY=$(tty)
echo "$CURRENT_TTY" > {shlex.quote(tty_file)}
printf "\\033]0;{shlex.quote(window_title)}\\007"

echo "========================================================"
echo " NOAA GSL Verification Studio"
echo " HPC Authentication for [{shlex.quote(target.upper())}]"
echo "========================================================"
echo ""

SUGGESTED_USER={shlex.quote(suggested_user)}
if [ -n "$SUGGESTED_USER" ]; then
    export HPC_{target_env_suffix}_USER="$SUGGESTED_USER"
    export HPC_USER="$SUGGESTED_USER"
fi

{shlex.quote(hpc_bin)} --start {shlex.quote(target)}
STATUS=$?
echo "$STATUS" > {shlex.quote(status_file)}

if [ $STATUS -eq 0 ]; then
    echo ""
    echo "========================================================"
    echo " [SUCCESS] Connection established! Closing window..."
    echo "========================================================"
    sleep 0.5
    osascript <<EOF >/dev/null 2>&1 &
tell application "Terminal"
    repeat with w in windows
        repeat with t in tabs of w
            if tty of t is "$CURRENT_TTY" then
                if (count of tabs of w) > 1 then
                    close t saving no
                else
                    close w saving no
                end if
                exit repeat
            end if
        end repeat
    end repeat
end tell
EOF
    exit 0
else
    echo ""
    echo "========================================================"
    echo " [ERROR] Authentication failed (exit code: $STATUS)."
    echo "========================================================"
    read -p "Press [Enter] to close..."
    exit $STATUS
fi
"""
    with open(cmd_file, "w", encoding="utf-8") as f:
        f.write(script_content)
    os.chmod(cmd_file, 0o700)

    try:
        subprocess.run(["open", "-a", "Terminal", cmd_file])

        t0 = time.time()
        while time.time() - t0 < 60:
            if os.path.exists(status_file):
                try:
                    with open(status_file, "r") as sf:
                        st = sf.read().strip()
                        if st == "0":
                            print(f"[HPC] >>> Successfully connected to {target}!", flush=True)
                            close_terminal_by_tty(tty_file)
                            return True
                        elif st != "":
                            print(f"[HPC] Authentication failed with status {st}.", flush=True)
                            close_terminal_by_tty(tty_file)
                            return False
                except Exception:
                    pass
            if is_hpc_online(target):
                print(f"[HPC] >>> Detected {target} is now ONLINE!", flush=True)
                close_terminal_by_tty(tty_file)
                return True
            time.sleep(1)

        print(f"[HPC] Authentication timed out after 60s for {target}.", flush=True)
        return False
    finally:
        close_terminal_by_tty(tty_file)
        if os.path.exists(auth_tempdir):
            try:
                shutil.rmtree(auth_tempdir)
            except Exception:
                pass


def fetch_remote_radar_images(models, cref_dir, radar_cases=None, force_fetch=False, on_progress=None):
    """
    Automated synoptic radar comparison image retrieval:
    Checks if evaluated models have remote HPC (gaea, hercules, etc.) or local paths specified.
    Downloads cref_full_sfc_f*.png for severe convective case cycles into cref_dir with local caching.
    """
    cases = radar_cases if radar_cases is not None else DEFAULT_RADAR_CASES
    if not cases:
        if on_progress: on_progress(100)
        return

    os.makedirs(cref_dir, exist_ok=True)
    hpc_bin = get_hpc_run_executable()

    models_with_path = [m for m in models if m.get("path")]
    if not models_with_path:
        if on_progress: on_progress(100)
        return

    total_radar_items = len(models_with_path) * len(cases)
    current_radar_idx = 0
    if on_progress: on_progress(0)

    for m in models_with_path:
        base_path = m.get("path")
        short_name = m.get("name")
        cluster = normalize_hpc_cluster(m.get("cluster"), base_path)

        is_remote = cluster in [
            "gaea", "hercules", "orion", "hera", "ursa",
            "gaea-c5", "gaea-c6", "gaea-c7",
            "gaeac5", "gaeac6", "gaeac7", "c5", "c6", "c7"
        ]

        # Pre-check which cases actually need downloading vs. already in local cache
        needed_cases = []
        for orig_idx, cdate, cycle, fhr, case_title in cases:
            fhr_str = f"{fhr:03d}"
            fhr_dir = f"{cdate}{cycle}{fhr:02d}"
            local_filename = f"{short_name}_{cdate}_{cycle}z_f{fhr:02d}.png"
            local_path = os.path.join(cref_dir, local_filename)
            if force_fetch or not is_valid_image(local_path):
                needed_cases.append((orig_idx, cdate, cycle, fhr, case_title, fhr_str, fhr_dir, local_filename, local_path))
            else:
                current_radar_idx += 1
                if on_progress:
                    on_progress(int(100 * (current_radar_idx / total_radar_items)))
                print(f"  [CACHE HIT] Using cached: {local_filename}", flush=True)

        if not needed_cases:
            print(f"\n>>> Checking synoptic radar images for [{short_name}] on {cluster or 'local'}...", flush=True)
            print(f"  [ALL CACHED] All {len(cases)} radar images for [{short_name}] already available in local cache.", flush=True)
            continue

        print(f"\n>>> Checking synoptic radar images for [{short_name}] on {cluster or 'local'} ({len(needed_cases)} to retrieve)...", flush=True)

        if is_remote:
            if not is_hpc_online(cluster):
                conn_ok = ensure_hpc_connection(cluster, models=models)
                if not conn_ok or not is_hpc_online(cluster):
                    print(f"  [NOTICE] HPC cluster '{cluster}' is offline. Skipping remote radar retrieval for [{short_name}].", flush=True)
                    current_radar_idx += len(needed_cases)
                    if on_progress:
                        on_progress(int(100 * (current_radar_idx / total_radar_items)))
                    continue

        env = os.environ.copy()
        env["HPC_NON_INTERACTIVE"] = "1"

        for orig_idx, cdate, cycle, fhr, case_title, fhr_str, fhr_dir, local_filename, local_path in needed_cases:
            current_radar_idx += 1
            if on_progress:
                on_progress(int(100 * (current_radar_idx / total_radar_items)))
            candidates = [
                f"{base_path}/com/rrfs/v2.1.4/rrfs.{cdate}/{cycle}/graphics/det/full/{fhr_dir}/cref_full_sfc_f{fhr_str}.png",
                f"{base_path}/rrfs.{cdate}/{cycle}/graphics/det/full/{fhr_dir}/cref_full_sfc_f{fhr_str}.png",
                f"{base_path}/graphics/det/full/{fhr_dir}/cref_full_sfc_f{fhr_str}.png",
                f"{base_path}/{fhr_dir}/cref_full_sfc_f{fhr_str}.png",
                f"{base_path}/{cdate}{cycle}/cref_full_sfc_f{fhr_str}.png",
                f"{base_path}/cref_full_sfc_f{fhr_str}.png",
            ]

            if is_remote:
                target_cluster = cluster
                cands_clean = [reject_pty_control_chars(c) for c in candidates]
                downloaded = False
                try:
                    for cand in cands_clean:
                        try:
                            res = subprocess.run(
                                [hpc_bin, target_cluster, "download", cand, local_path],
                                capture_output=True,
                                text=True,
                                timeout=15,
                                env=env
                            )
                            if res.returncode == 0 and os.path.exists(local_path) and os.path.getsize(local_path) > 0:
                                sz = os.path.getsize(local_path)
                                print(f"  [CACHE MISS] Downloaded fresh image: {local_filename} from {target_cluster} ({sz} bytes)", flush=True)
                                downloaded = True
                                break
                        except subprocess.TimeoutExpired:
                            pass

                    if not downloaded:
                        print(f"  [NOTICE] No remote radar image found for {local_filename} on {target_cluster}:{base_path}", flush=True)
                        if not is_hpc_online(target_cluster):
                            print(f"  [WARNING] HPC cluster '{target_cluster}' connection dropped. Prompting to reconnect...", flush=True)
                            if ensure_hpc_connection(target_cluster, models=models):
                                print(f"  [SUCCESS] Reconnected to '{target_cluster}'! Retrying {local_filename}...", flush=True)
                                for cand in cands_clean:
                                    try:
                                        res_retry = subprocess.run(
                                            [hpc_bin, target_cluster, "download", cand, local_path],
                                            capture_output=True,
                                            text=True,
                                            timeout=15,
                                            env=env
                                        )
                                        if res_retry.returncode == 0 and os.path.exists(local_path) and os.path.getsize(local_path) > 0:
                                            sz = os.path.getsize(local_path)
                                            print(f"  [CACHE MISS] Downloaded fresh image: {local_filename} from {target_cluster} ({sz} bytes)", flush=True)
                                            downloaded = True
                                            break
                                    except Exception:
                                        pass
                            else:
                                print(f"  [NOTICE] Re-authentication cancelled or failed. Skipping remaining remote checks for [{short_name}].", flush=True)
                                break
                except Exception as e:
                    print(f"  [WARNING] Error fetching remote {local_filename}: {e}", flush=True)
            else:
                found = False
                for cand in candidates:
                    if os.path.exists(cand):
                        shutil.copyfile(cand, local_path)
                        print(f"  [CACHE MISS] Copied fresh image: {local_filename} from local {cand}", flush=True)
                        found = True
                        break
                if not found:
                    print(f"  [NOTICE] Local image not found: {local_filename}", flush=True)

    if on_progress:
        on_progress(100)





async def launch_browser(p, headless=True):
    # Try system Google Chrome if available
    if os.path.exists('/Applications/Google Chrome.app'):
        try:
            return await p.chromium.launch(headless=headless, channel='chrome', args=['--disable-blink-features=AutomationControlled'])
        except Exception:
            pass
    # Fallback to bundled or standard Playwright Chromium
    return await p.chromium.launch(headless=headless, args=['--disable-blink-features=AutomationControlled'])

async def fetch_mrms_radar_images(p, cref_dir, radar_cases=None, force_fetch=False, on_progress=None):
    """
    Automated MRMS observation retrieval directly from NOAA NSSL MRMS Product Viewer:
    https://mrms.nssl.noaa.gov/qvs/product_viewer/
    Calculates valid verification time (cycle_time + fhr) and captures
    the high-resolution composite reflectivity map with dBZ legend into cref_dir.
    """
    cases = radar_cases if radar_cases is not None else DEFAULT_RADAR_CASES
    if not cases:
        if on_progress: on_progress(100)
        return

    os.makedirs(cref_dir, exist_ok=True)
    if on_progress: on_progress(0)
    print("\n>>> Checking MRMS observations directly from NOAA NSSL Product Viewer...", flush=True)

    needed = []
    total_cases = len(cases)
    current_mrms_idx = 0
    for orig_idx, cdate, cycle, fhr, case_title in cases:
        mrms_fn = f"mrms_{cdate}_{cycle}z_f{fhr:02d}.png"
        target_path = os.path.join(cref_dir, mrms_fn)
        if force_fetch or not is_valid_image(target_path):
            needed.append((orig_idx, cdate, cycle, fhr, case_title, target_path, mrms_fn))
        else:
            current_mrms_idx += 1
            if on_progress:
                on_progress(int(100 * (current_mrms_idx / total_cases)))
            print(f"  [CACHE HIT] Using cached: {mrms_fn}", flush=True)

    if not needed:
        print("  [ALL CACHED] All MRMS observation radar images already available locally.", flush=True)
        if on_progress: on_progress(100)
        return

    print(f"  [FETCHING] Downloading {len(needed)} MRMS observation images from NSSL Product Viewer...", flush=True)
    browser = None
    try:
        browser = await launch_browser(p)
        page = await browser.new_page(viewport={"width": 1600, "height": 1000}, device_scale_factor=1)

        for orig_idx, cdate, cycle, fhr, case_title, target_path, mrms_fn in needed:
            current_mrms_idx += 1
            if on_progress:
                on_progress(int(100 * (current_mrms_idx / total_cases)))
            try:
                init_dt = datetime.datetime.strptime(f"{cdate}{cycle}", "%Y%m%d%H")
                valid_dt = init_dt + datetime.timedelta(hours=fhr)
                y, m, d, h, mn = valid_dt.year, valid_dt.month, valid_dt.day, valid_dt.hour, valid_dt.minute

                url = (
                    f"https://mrms.nssl.noaa.gov/qvs/product_viewer/index.php?"
                    f"web_exec_mode=run&menu=menu_config.txt&year={y}&month={m}&day={d}&hour={h}&minute={mn}"
                    f"&time_mode=static&zoom=4&clon=-98&clat=38&base=0&overlays=1&mping_mode=0&product_type=crefls&product=CREF"
                )
                print(f"  • Fetching MRMS for {case_title} (Valid: {valid_dt.strftime('%Y-%m-%d %H:%M UTC')})...", flush=True)

                await page.goto(url, wait_until="domcontentloaded", timeout=25000)
                await page.wait_for_selector("#base_map", timeout=15000)
                await asyncio.sleep(2.0)  # Allow canvas drawing to finish

                base_map = await page.query_selector("#base_map")
                side_legend = await page.query_selector("#side_legend")
                if base_map and side_legend:
                    b_box = await base_map.bounding_box()
                    l_box = await side_legend.bounding_box()
                    if b_box and l_box:
                        x = b_box["x"]
                        y = b_box["y"]
                        w = (l_box["x"] + l_box["width"]) - x
                        h = max(b_box["height"], l_box["height"]) + 3
                        tmp_target_path = target_path + f".tmp.{os.getpid()}.png"
                        await page.screenshot(path=tmp_target_path, clip={"x": x, "y": y, "width": w, "height": h})
                        os.replace(tmp_target_path, target_path)
                        print(f"    [CACHE MISS] Downloaded fresh MRMS observation: {mrms_fn} ({os.path.getsize(target_path)} bytes)", flush=True)
                        continue

                print(f"    [WARNING] Could not locate map/legend selectors for {mrms_fn}", flush=True)
            except Exception as e:
                print(f"    [WARNING] Error fetching MRMS {mrms_fn}: {e}", flush=True)
    except Exception as ex:
        print(f"  [ERROR] MRMS browser automation failed: {ex}", flush=True)
    finally:
        if browser:
            await browser.close()

    if on_progress:
        on_progress(100)




async def check_available_models(page, category_name, models, variable=None, database=None):
    """
    Universal model availability pre-checker across all MATS applications.
    Returns (valid_models, missing_models). Gracefully logs status.
    """
    try:
        if database:
            await page.evaluate(f'''() => {{
                if (window.matsParamUtils) {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('database', '{database}');
                }}
            }}''')
            await page.wait_for_timeout(1000)

        if variable:
            await page.evaluate(f'''() => {{
                if (window.matsParamUtils) {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{variable}');
                }}
            }}''')
            await page.wait_for_timeout(1000)

        avail = await page.evaluate('''() => {
            const select = document.querySelector('select[name="data-source"]');
            return select ? Array.from(select.options).map(o => o.value) : [];
        }''')

        valid = [m for m in models if m["id"] in avail]
        missing = [m for m in models if m["id"] not in avail]

        if not valid:
            print(f"  [{category_name.upper()} NOTICE] None of the requested models exist in MATS {category_name} database. Gracefully skipping.")
            return [], missing

        if missing:
            print(f"  [{category_name.upper()} NOTICE] Partially available: {len(valid)} of {len(models)} models found ({', '.join([m['name'] for m in valid])}). Missing: {', '.join([m['name'] for m in missing])}.")
        else:
            print(f"  [{category_name.upper()}] All {len(valid)} models available in database.")

        return valid, missing
    except Exception as e:
        print(f"  [{category_name.upper()} WARNING] Could not verify model availability: {e}. Proceeding with all models.")
        return models, []


async def export_styled_chart(page, outpath, color_map, width=1761, height=857, models=None):
    """
    Applies custom styling to Plotly charts inside MATS:
    - Same experiment -> Same color
    - RMSE / CSI -> Solid line, renamed to clean short experiment name
    - BIAS -> Dash line, hidden from legend if main metric is present
    - Non-intrusive compact legend layout
    - Exports directly via Plotly.toImage (no web UI)
    """
    # 1. Restyle curves via Plotly
    meta = {
        "colorMap": color_map or {},
        "models": [{"id": m.get("id"), "name": m.get("name"), "color": m.get("color")} for m in (models or [])]
    }
    await page.evaluate(f'''(meta) => {{
        const colorMap = meta.colorMap || {{}};
        const models = meta.models || [];
        const plots = Array.from(document.querySelectorAll('.js-plotly-plot'));
        const gd = plots.find(el => el.offsetWidth > 200 && el.offsetHeight > 200);
        if (!gd || !gd.data) return;
        
        // Detect if chart has both main metric (RMSE/CSI) and BIAS
        const hasMainMetric = gd.data.some(t => {{
            const n = (t.name || "").toLowerCase();
            return !n.includes("bias") && n !== "x = 0" && n !== "y = 0";
        }});
        
        gd.data.forEach((trace) => {{
            if (!trace.name || trace.name === "x = 0" || trace.name === "y = 0") return;
            
            let matchedModel = null;
            for (const m of models) {{
                if (trace.name.includes(m.id) || trace.name.includes(m.name)) {{
                    matchedModel = m;
                    break;
                }}
            }}
            
            let matchedColor = matchedModel ? matchedModel.color : null;
            if (!matchedColor) {{
                for (const [key, color] of Object.entries(colorMap)) {{
                    if (trace.name.includes(key)) {{
                        matchedColor = color;
                        break;
                    }}
                }}
            }}
            
            const isBias = trace.name.toLowerCase().includes("bias");
            const color = matchedColor || (trace.line && trace.line.color) || "#2563eb";
            const shortName = matchedModel ? matchedModel.name : trace.name.split(':')[0];
            
            if (trace.line) {{
                trace.line.color = color;
                trace.line.dash = isBias ? "dash" : "solid";
                trace.line.width = 3;
            }}
            if (trace.marker) trace.marker.color = color;
            if (trace.error_y) trace.error_y.color = color;
            if (trace.error_x) trace.error_x.color = color;
            
            // Clean up legend:
            if (isBias && hasMainMetric) {{
                trace.showlegend = false; // Hide BIAS duplicate when main metric exists
            }} else {{
                trace.name = shortName;   // Use clean short experiment name
                trace.showlegend = true;
            }}
        }});
        
        if (gd.layout && gd.layout.legend) {{
            gd.layout.legend.font = {{ size: 12, color: '#1e293b' }};
            gd.layout.legend.bgcolor = 'rgba(255, 255, 255, 0.85)';
            gd.layout.legend.bordercolor = 'rgba(203, 213, 225, 0.8)';
            gd.layout.legend.borderwidth = 1;
        }}
        
        Plotly.redraw(gd);
    }}''', meta)
    await page.wait_for_timeout(800)
    
    # 2. Pure canvas export
    img_data = await page.evaluate(f'''async () => {{
        const plots = Array.from(document.querySelectorAll('.js-plotly-plot'));
        const gd = plots.find(el => el.offsetWidth > 200 && el.offsetHeight > 200);
        if (window.Plotly && gd) {{
            return await Plotly.toImage(gd, {{
                format: 'png',
                width: {width},
                height: {height}
            }});
        }}
        return null;
    }}''')
    if img_data and img_data.startswith('data:image/png;base64,'):
        data = base64.b64decode(img_data.split(',', 1)[1])
        atomic_write_bytes(outpath, data)
        print(f"  [CACHE MISS] Exported fresh chart: {os.path.basename(outpath)} ({len(data)} bytes)")
        return True
    print(f"  [FAILED] Export failed for {os.path.basename(outpath)}")
    return False


async def fetch_surface_dieoff(p, models, sfc_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [1/8] Fetching Surface Dieoff Charts (Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/surface", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Surface Dieoff", models)
    except Exception as e:
        print(f"  [SURFACE NOTICE] Failed to connect to MATS Surface: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    vars_list = [
        ("Temperature at 2m (°C)", "t2m"),
        ("Dewpoint at 2m (°C)", "td2m"),
        ("Wind Speed at 10m (m/s)", "wind10m")
    ]
    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["surface"].get(reg, "All HRRR domain")
        for var, vtag in vars_list:
            fn = f"surface_dieoff_{vtag}{reg_tag}.png"
            outpath = os.path.join(out_dir, fn)
            if not force_fetch and is_valid_image(outpath):
                print(f"  [CACHE HIT] Using cached: {fn}")
                continue
                
            b = await launch_browser(p)
            c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
            pg = await c.new_page()
            try:
                print(f"  Processing Surface Dieoff ({reg}): {var}...")
                await pg.goto("https://gsl.noaa.gov/mats/surface", timeout=45000)
                await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                await pg.select_option("#plotTypes-selector", value="Dieoff")
                await pg.wait_for_timeout(600)
                
                # RMSE curves
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'RMSE');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                        matsParamUtils.setInputValueForParamAndTriggerChange('truth', 'METAR');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                # Bias curves
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Model - Obs)');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                        matsParamUtils.setInputValueForParamAndTriggerChange('truth', 'METAR');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                await pg.click("#plotMatched")
                await pg.wait_for_timeout(8000)
                await export_styled_chart(pg, outpath, color_map, models=valid_models, width=1761, height=857)
            except Exception as e:
                print(f"  [WARNING] Failed to fetch Surface Dieoff ({var}, {reg}): {e}. Skipping.")
            finally:
                await b.close()


async def fetch_surface_timeseries(p, models, sfc_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [2/8] Fetching Surface Time Series (1h, 3h, 12h, 24h; Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/surface", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Surface TimeSeries", models)
    except Exception as e:
        print(f"  [SURFACE TS NOTICE] Failed to connect to MATS Surface: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    configs = [
        ("1", "RMSE", "Temperature at 2m (°C)", "sfc_ts_1h_rmse_t2m.png"),
        ("1", "RMSE", "Dewpoint at 2m (°C)", "sfc_ts_1h_rmse_td2m.png"),
        ("1", "RMSE", "Wind Speed at 10m (m/s)", "sfc_ts_1h_rmse_wind10m.png"),
        ("1", "Bias (Model - Obs)", "Temperature at 2m (°C)", "sfc_ts_1h_bias_t2m.png"),
        ("1", "Bias (Model - Obs)", "Dewpoint at 2m (°C)", "sfc_ts_1h_bias_td2m.png"),
        ("1", "Bias (Model - Obs)", "Wind Speed at 10m (m/s)", "sfc_ts_1h_bias_wind10m.png"),
        ("3", "RMSE", "Temperature at 2m (°C)", "sfc_ts_3h_rmse_t2m.png"),
        ("3", "RMSE", "Dewpoint at 2m (°C)", "sfc_ts_3h_rmse_td2m.png"),
        ("3", "RMSE", "Wind Speed at 10m (m/s)", "sfc_ts_3h_rmse_wind10m.png"),
        ("3", "Bias (Model - Obs)", "Temperature at 2m (°C)", "sfc_ts_3h_bias_t2m.png"),
        ("3", "Bias (Model - Obs)", "Dewpoint at 2m (°C)", "sfc_ts_3h_bias_td2m.png"),
        ("3", "Bias (Model - Obs)", "Wind Speed at 10m (m/s)", "sfc_ts_3h_bias_wind10m.png"),
        ("12", "RMSE", "Temperature at 2m (°C)", "sfc_ts_12h_rmse_t2m.png"),
        ("12", "RMSE", "Dewpoint at 2m (°C)", "sfc_ts_12h_rmse_td2m.png"),
        ("12", "RMSE", "Wind Speed at 10m (m/s)", "sfc_ts_12h_rmse_wind10m.png"),
        ("12", "Bias (Model - Obs)", "Temperature at 2m (°C)", "sfc_ts_12h_bias_t2m.png"),
        ("12", "Bias (Model - Obs)", "Dewpoint at 2m (°C)", "sfc_ts_12h_bias_td2m.png"),
        ("12", "Bias (Model - Obs)", "Wind Speed at 10m (m/s)", "sfc_ts_12h_bias_wind10m.png"),
        ("24", "RMSE", "Temperature at 2m (°C)", "sfc_ts_24h_rmse_t2m.png"),
        ("24", "RMSE", "Dewpoint at 2m (°C)", "sfc_ts_24h_rmse_td2m.png"),
        ("24", "RMSE", "Wind Speed at 10m (m/s)", "sfc_ts_24h_rmse_wind10m.png"),
        ("24", "Bias (Model - Obs)", "Temperature at 2m (°C)", "sfc_ts_24h_bias_t2m.png"),
        ("24", "Bias (Model - Obs)", "Dewpoint at 2m (°C)", "sfc_ts_24h_bias_td2m.png"),
        ("24", "Bias (Model - Obs)", "Wind Speed at 10m (m/s)", "sfc_ts_24h_bias_wind10m.png"),
    ]
    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["surface"].get(reg, "All HRRR domain")
        for fhr, stat, var, base_fn in configs:
            fn = base_fn.replace(".png", f"{reg_tag}.png")
            outpath = os.path.join(out_dir, fn)
            if not force_fetch and is_valid_image(outpath):
                print(f"  [CACHE HIT] Using cached: {fn}")
                continue
                
            b = await launch_browser(p)
            c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
            pg = await c.new_page()
            try:
                print(f"  Processing Surface TS ({reg}): {fn} (fhr={fhr}, {stat}, {var})...")
                await pg.goto("https://gsl.noaa.gov/mats/surface", timeout=45000)
                await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                await pg.select_option("#plotTypes-selector", value="TimeSeries")
                await pg.wait_for_timeout(600)
                
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', '{stat}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('truth', 'METAR');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                await pg.click("#plotMatched")
                await pg.wait_for_timeout(8000)
                await export_styled_chart(pg, outpath, color_map, models=valid_models, width=1761, height=857)
            except Exception as e:
                print(f"  [WARNING] Failed to fetch Surface TS ({fn}, {reg}): {e}. Skipping.")
            finally:
                await b.close()


async def fetch_upperair_dieoff(p, models, ua_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [3/8] Fetching Upper Air Dieoff against RAOBs (Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Upper Air RAOBs", models)
    except Exception as e:
        print(f"  [UPPER AIR NOTICE] Failed to connect to MATS UpperAir: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    vars_list = [
        ("Temperature", "ua_dieoff_temp.png"),
        ("RHobT", "ua_dieoff_rh.png"),
        ("Wind", "ua_dieoff_wind.png")
    ]
    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["upperair"].get(reg, "HRRR domain")
        for var, base_fn in vars_list:
            fn = base_fn.replace(".png", f"{reg_tag}.png")
            outpath = os.path.join(out_dir, fn)
            if not force_fetch and is_valid_image(outpath):
                print(f"  [CACHE HIT] Using cached: {fn}")
                continue
                
            b = await launch_browser(p)
            c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
            pg = await c.new_page()
            try:
                print(f"  Processing UA Dieoff ({reg}): {var}...")
                await pg.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
                await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                await pg.select_option("#plotTypes-selector", value="Dieoff")
                await pg.wait_for_timeout(600)
                
                # RMSE
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'RMSE');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                # Bias
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Model - Obs)');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                await pg.click("#plotMatched")
                await pg.wait_for_timeout(8000)
                await export_styled_chart(pg, outpath, color_map, models=valid_models, width=1761, height=857)
            except Exception as e:
                print(f"  [WARNING] Failed to fetch UA Dieoff ({var}, {reg}): {e}. Skipping.")
            finally:
                await b.close()


async def fetch_upperair_profiles(p, models, ua_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [4/8] Fetching Upper Air Profiles (1h, 3h, 6h, 12h; Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Upper Air Profiles", models)
    except Exception as e:
        print(f"  [UPPER AIR PROFILE NOTICE] Failed to connect to MATS UpperAir: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    configs = [
        ("1", "Temperature", "ua_prof_1h_temp.png"),
        ("1", "Wind", "ua_prof_1h_wind.png"),
        ("1", "RHobT", "ua_prof_1h_rh.png"),
        ("3", "Temperature", "ua_prof_3h_temp.png"),
        ("3", "Wind", "ua_prof_3h_wind.png"),
        ("3", "RHobT", "ua_prof_3h_rh.png"),
        ("6", "Temperature", "ua_prof_6h_temp.png"),
        ("6", "Wind", "ua_prof_6h_wind.png"),
        ("6", "RHobT", "ua_prof_6h_rh.png"),
        ("12", "Temperature", "ua_prof_12h_temp.png"),
        ("12", "Wind", "ua_prof_12h_wind.png"),
        ("12", "RHobT", "ua_prof_12h_rh.png"),
    ]
    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["upperair"].get(reg, "HRRR domain")
        for fhr, var, base_fn in configs:
            fn = base_fn.replace(".png", f"{reg_tag}.png")
            outpath = os.path.join(out_dir, fn)
            if not force_fetch and is_valid_image(outpath):
                print(f"  [CACHE HIT] Using cached: {fn}")
                continue
                
            b = await launch_browser(p)
            c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
            pg = await c.new_page()
            try:
                print(f"  Processing UA Profile ({reg}): {fn} (fhr={fhr}, {var})...")
                await pg.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
                await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                await pg.select_option("#plotTypes-selector", value="Profile")
                await pg.wait_for_timeout(600)
                
                # RMSE
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'RMSE');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                # Bias
                for m in valid_models:
                    await pg.evaluate(f'''() => {{
                        matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                        matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Model - Obs)');
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                        matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                    }}''')
                    await pg.wait_for_timeout(150)
                    await pg.click("#add")
                    await pg.wait_for_timeout(150)
                    
                await pg.click("#plotMatched")
                await pg.wait_for_timeout(8000)
                await export_styled_chart(pg, outpath, color_map, models=valid_models, width=900, height=900)
            except Exception as e:
                print(f"  [WARNING] Failed to fetch UA Profile ({fn}, {reg}): {e}. Skipping.")
            finally:
                await b.close()


async def fetch_upperair_amdar(p, models, ua_dates, color_map, out_dir, force_fetch=False):
    print("\n>>> [5/8] Checking Upper Air AMDAR Verification (Dieoff & Profiles)...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
        await page.wait_for_selector("#plotTypes-selector", timeout=25000)
        valid_models, _ = await check_available_models(page, "AMDAR", models, database='AMDAR (Traditional)')
    except Exception as e:
        print(f"  [AMDAR NOTICE] Could not query AMDAR database: {e}. Skipping AMDAR.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    # 1. AMDAR Dieoff
    vars_list = [
        ("Temperature", "amdar_dieoff_temp.png"),
        ("Wind", "amdar_dieoff_wind.png"),
        ("RHobT", "amdar_dieoff_rh.png")
    ]
    for var, fn in vars_list:
        outpath = os.path.join(out_dir, fn)
        if not force_fetch and is_valid_image(outpath):
            print(f"  [CACHE HIT] Using cached: {fn}")
            continue
            
        b = await launch_browser(p)
        c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
        pg = await c.new_page()
        try:
            print(f"  Processing AMDAR Dieoff: {var}...")
            await pg.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
            await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
            await pg.select_option("#plotTypes-selector", value="Dieoff")
            await pg.wait_for_timeout(600)
            
            await pg.evaluate('''() => {
                matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
            }''')
            await pg.wait_for_timeout(1000)
            
            for m in valid_models:
                await pg.evaluate(f'''() => {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('region', 'Full model domain');
                    matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'RMSE');
                    matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                    matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                }}''')
                await pg.wait_for_timeout(150)
                await pg.click("#add")
                await pg.wait_for_timeout(150)
                
            for m in valid_models:
                await pg.evaluate(f'''() => {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('region', 'Full model domain');
                    matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Model - Obs)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                    matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                }}''')
                await pg.wait_for_timeout(150)
                await pg.click("#add")
                await pg.wait_for_timeout(150)
                
            await pg.click("#plotMatched")
            await pg.wait_for_timeout(8000)
            await export_styled_chart(pg, outpath, color_map, models=valid_models)
        except Exception as e:
            print(f"  [WARNING] Failed to fetch AMDAR Dieoff ({var}): {e}. Skipping.")
        finally:
            await b.close()

    # 2. AMDAR Profiles (1h, 3h, 6h, 12h)
    amdar_prof_configs = [
        ("1", "Temperature", "amdar_prof_1h_temp.png"),
        ("1", "Wind", "amdar_prof_1h_wind.png"),
        ("1", "RHobT", "amdar_prof_1h_rh.png"),
        ("3", "Temperature", "amdar_prof_3h_temp.png"),
        ("3", "Wind", "amdar_prof_3h_wind.png"),
        ("3", "RHobT", "amdar_prof_3h_rh.png"),
        ("6", "Temperature", "amdar_prof_6h_temp.png"),
        ("6", "Wind", "amdar_prof_6h_wind.png"),
        ("6", "RHobT", "amdar_prof_6h_rh.png"),
        ("12", "Temperature", "amdar_prof_12h_temp.png"),
        ("12", "Wind", "amdar_prof_12h_wind.png"),
        ("12", "RHobT", "amdar_prof_12h_rh.png"),
    ]
    for fhr, var, fn in amdar_prof_configs:
        outpath = os.path.join(out_dir, fn)
        if not force_fetch and is_valid_image(outpath):
            print(f"  [CACHE HIT] Using cached: {fn}")
            continue

        b = await launch_browser(p)
        c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
        pg = await c.new_page()
        try:
            print(f"  Processing AMDAR Profile: {fn} (fhr={fhr}, {var})...")
            await pg.goto("https://gsl.noaa.gov/mats/upperair", timeout=45000)
            await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
            await pg.select_option("#plotTypes-selector", value="Profile")
            await pg.wait_for_timeout(600)
            await pg.evaluate('''() => {
                matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
            }''')
            await pg.wait_for_timeout(1000)

            # RMSE
            for m in valid_models:
                await pg.evaluate(f'''() => {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('region', 'Full model domain');
                    matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'RMSE');
                    matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                }}''')
                await pg.wait_for_timeout(150)
                await pg.click("#add")
                await pg.wait_for_timeout(150)

            # Bias
            for m in valid_models:
                await pg.evaluate(f'''() => {{
                    matsParamUtils.setInputValueForParamAndTriggerChange('database', 'AMDAR (Traditional)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('region', 'Full model domain');
                    matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Model - Obs)');
                    matsParamUtils.setInputValueForParamAndTriggerChange('variable', '{var}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                    matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{ua_dates}');
                }}''')
                await pg.wait_for_timeout(150)
                await pg.click("#add")
                await pg.wait_for_timeout(150)

            await pg.click("#plotMatched")
            await pg.wait_for_timeout(8000)
            await export_styled_chart(pg, outpath, color_map, models=valid_models, width=900, height=900)
        except Exception as e:
            print(f"  [WARNING] Failed to fetch AMDAR Profile ({fn}): {e}. Skipping.")
        finally:
            await b.close()


async def fetch_reflectivity(p, models, sfc_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [6/8] Fetching Composite Reflectivity Verification (Dieoff & 1h/3h TS for 15, 25, 35 dBZ; Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/radar", timeout=45000)
        await page.wait_for_timeout(3000)
        valid_models, _ = await check_available_models(page, "Radar Reflectivity", models, variable="Composite Reflectivity")
    except Exception as e:
        print(f"  [RADAR NOTICE] Failed to connect to MATS Radar: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    thresholds = [
        ("15 (reflectivity >= 15 dBZ)", "15dbz", "15 dBZ"),
        ("25 (reflectivity >= 25 dBZ)", "25dbz", "25 dBZ"),
        ("35 (reflectivity >= 35 dBZ)", "35dbz", "35 dBZ"),
    ]

    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["radar"].get(reg, "Continental US")

        # 1. Dieoff: CSI & Bias
        for thresh_str, tag, lbl in thresholds:
            fn_csi = f"ref_dieoff_csi_{tag}{reg_tag}.png"
            fn_bias = f"ref_dieoff_bias_{tag}{reg_tag}.png"
            out_csi = os.path.join(out_dir, fn_csi)
            out_bias = os.path.join(out_dir, fn_bias)

            # Dieoff CSI
            if force_fetch or not is_valid_image(out_csi):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Radar CSI Dieoff ({reg}): {lbl}...")
                    await pg.goto("https://gsl.noaa.gov/mats/radar", timeout=45000)
                    await pg.wait_for_timeout(2500)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(600)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('scale', '3 km grid');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Radar CSI Dieoff ({lbl}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_csi}")

            # Dieoff Bias
            if force_fetch or not is_valid_image(out_bias):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Radar Bias Dieoff ({reg}): {lbl}...")
                    await pg.goto("https://gsl.noaa.gov/mats/radar", timeout=45000)
                    await pg.wait_for_timeout(2500)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(600)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('scale', '3 km grid');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Radar Bias Dieoff ({lbl}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_bias}")

        # 2. Time Series: 1-h and 3-h fcst (CSI & Bias)
        for fhr in ["1", "3"]:
            for thresh_str, tag, lbl in thresholds:
                fn_csi = f"ref_ts_{fhr}h_csi_{tag}{reg_tag}.png"
                fn_bias = f"ref_ts_{fhr}h_bias_{tag}{reg_tag}.png"
                out_csi = os.path.join(out_dir, fn_csi)
                out_bias = os.path.join(out_dir, fn_bias)

                # TS CSI
                if force_fetch or not is_valid_image(out_csi):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Radar {fhr}-h TS CSI ({reg}): {lbl}...")
                        await pg.goto("https://gsl.noaa.gov/mats/radar", timeout=45000)
                        await pg.wait_for_timeout(2500)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('scale', '3 km grid');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Radar {fhr}-h TS CSI ({lbl}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_csi}")

                # TS Bias
                if force_fetch or not is_valid_image(out_bias):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Radar {fhr}-h TS Bias ({reg}): {lbl}...")
                        await pg.goto("https://gsl.noaa.gov/mats/radar", timeout=45000)
                        await pg.wait_for_timeout(2500)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('scale', '3 km grid');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Composite Reflectivity');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Radar {fhr}-h TS Bias ({lbl}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_bias}")


async def fetch_ceiling(p, models, sfc_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [7/8] Fetching Ceiling Verification (Dieoff & 1h/3h TS; Regions: {', '.join(regions)})...")
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Ceiling", models, variable='Ceiling')
    except Exception as e:
        print(f"  [CEILING NOTICE] Failed to query MATS Ceiling: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    thresholds = [
        ("500 (ceiling <500 ft)", "500ft"),
        ("1000 (ceiling <1000 ft)", "1000ft"),
        ("3000 (ceiling <3000 ft)", "3000ft"),
    ]
    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["ceil_vis"].get(reg, "All HRRR domain")

        # 1. Dieoff
        for thresh_str, tag in thresholds:
            thresh_val = thresh_str
            fn_csi = f"ceil_dieoff_csi_{tag}{reg_tag}.png"
            fn_bias = f"ceil_dieoff_bias_{tag}{reg_tag}.png"
            out_csi = os.path.join(out_dir, fn_csi)
            out_bias = os.path.join(out_dir, fn_bias)
            
            if force_fetch or not is_valid_image(out_csi):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Ceiling CSI Dieoff ({reg}): {tag} ({thresh_val})...")
                    await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                    await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(800)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_val}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Ceiling CSI Dieoff ({tag}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_csi}")
                
            if force_fetch or not is_valid_image(out_bias):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Ceiling Bias Dieoff ({reg}): {tag} ({thresh_val})...")
                    await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                    await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(800)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_val}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Ceiling Bias Dieoff ({tag}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_bias}")

        # 2. Time Series: 1-h and 3-h fcst
        for fhr in ["1", "3"]:
            for thresh_str, tag in thresholds:
                fn_csi = f"ceil_ts_{fhr}h_csi_{tag}{reg_tag}.png"
                fn_bias = f"ceil_ts_{fhr}h_bias_{tag}{reg_tag}.png"
                out_csi = os.path.join(out_dir, fn_csi)
                out_bias = os.path.join(out_dir, fn_bias)

                if force_fetch or not is_valid_image(out_csi):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Ceiling {fhr}-h TS CSI ({reg}): {tag}...")
                        await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                        await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Ceiling {fhr}-h TS CSI ({tag}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_csi}")

                if force_fetch or not is_valid_image(out_bias):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Ceiling {fhr}-h TS Bias ({reg}): {tag}...")
                        await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                        await pg.wait_for_selector("#plotTypes-selector", timeout=25000)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Ceiling');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Ceiling {fhr}-h TS Bias ({tag}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_bias}")


async def fetch_visibility(p, models, sfc_dates, color_map, out_dir, regions=None, force_fetch=False):
    regions = regions or ["all"]
    print(f"\n>>> [8/8] Fetching Visibility Verification (Dieoff & 1h/3h TS; Regions: {', '.join(regions)})...")
    thresholds = [
        ("1 (visibility < 1 mi)", "1mi"),
        ("3 (visibility < 3 mi)", "3mi"),
        ("5 (visibility < 5 mi)", "5mi"),
    ]
    
    browser = await launch_browser(p)
    context = await browser.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
    page = await context.new_page()
    valid_models = []
    try:
        await page.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
        await page.wait_for_timeout(2500)
        valid_models, _ = await check_available_models(page, "Visibility", models, variable='Visibility')
    except Exception as e:
        print(f"  [VISIBILITY NOTICE] Failed to query MATS Visibility: {e}. Gracefully skipping.")
        await browser.close()
        return
    finally:
        await browser.close()

    if not valid_models:
        return

    for reg in regions:
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_label = REGION_MAP["ceil_vis"].get(reg, "All HRRR domain")

        # 1. Dieoff CSI & Bias
        for thresh_str, tag in thresholds:
            fn_csi = f"vis_dieoff_csi_{tag}{reg_tag}.png"
            fn_bias = f"vis_dieoff_bias_{tag}{reg_tag}.png"
            out_csi = os.path.join(out_dir, fn_csi)
            out_bias = os.path.join(out_dir, fn_bias)
            
            if force_fetch or not is_valid_image(out_csi):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Visibility CSI Dieoff ({reg}): {tag}...")
                    await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                    await pg.wait_for_timeout(2500)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(600)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Visibility CSI Dieoff ({tag}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_csi}")
                
            if force_fetch or not is_valid_image(out_bias):
                b = await launch_browser(p)
                c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                pg = await c.new_page()
                try:
                    print(f"  Processing Visibility Bias Dieoff ({reg}): {tag}...")
                    await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                    await pg.wait_for_timeout(2500)
                    await pg.select_option("#plotTypes-selector", value="Dieoff")
                    await pg.wait_for_timeout(600)
                    await pg.evaluate('''() => {
                        matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                    }''')
                    await pg.wait_for_timeout(600)
                    for m in valid_models:
                        await pg.evaluate(f'''() => {{
                            matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                            matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                            matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                            matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                            matsParamUtils.setInputValueForParamAndTriggerChange('dieoff-type', 'Dieoff');
                            matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                        }}''')
                        await pg.wait_for_timeout(150)
                        await pg.click("#add")
                        await pg.wait_for_timeout(150)
                    await pg.click("#plotMatched")
                    await pg.wait_for_timeout(8000)
                    await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                except Exception as e:
                    print(f"  [WARNING] Failed to fetch Visibility Bias Dieoff ({tag}, {reg}): {e}. Skipping.")
                finally:
                    await b.close()
            else:
                print(f"  [CACHE HIT] Using cached: {fn_bias}")
                
        # 2. Time Series: 1-h and 3-h fcst (CSI & Bias)
        for fhr in ["1", "3"]:
            for thresh_str, tag in thresholds:
                fn_csi = f"vis_ts_{fhr}h_csi_{tag}{reg_tag}.png"
                fn_bias = f"vis_ts_{fhr}h_bias_{tag}{reg_tag}.png"
                out_csi = os.path.join(out_dir, fn_csi)
                out_bias = os.path.join(out_dir, fn_bias)
                
                if force_fetch or not is_valid_image(out_csi):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Visibility {fhr}-h TS CSI ({reg}): {tag}...")
                        await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                        await pg.wait_for_timeout(2500)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'CSI (Critical Success Index)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_csi, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Visibility {fhr}-h TS CSI ({tag}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_csi}")
                    
                if force_fetch or not is_valid_image(out_bias):
                    b = await launch_browser(p)
                    c = await b.new_context(user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36")
                    pg = await c.new_page()
                    try:
                        print(f"  Processing Visibility {fhr}-h TS Bias ({reg}): {tag}...")
                        await pg.goto("https://gsl.noaa.gov/mats/ceil-vis", timeout=45000)
                        await pg.wait_for_timeout(2500)
                        await pg.select_option("#plotTypes-selector", value="TimeSeries")
                        await pg.wait_for_timeout(600)
                        await pg.evaluate('''() => {
                            matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                        }''')
                        await pg.wait_for_timeout(600)
                        for m in valid_models:
                            await pg.evaluate(f'''() => {{
                                matsParamUtils.setInputValueForParamAndTriggerChange('data-source', '{m["id"]}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region-type', 'Predefined region');
                                matsParamUtils.setInputValueForParamAndTriggerChange('region', '{reg_label}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('variable', 'Visibility');
                                matsParamUtils.setInputValueForParamAndTriggerChange('threshold', '{thresh_str}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('statistic', 'Bias (Forecast / Actual)');
                                matsParamUtils.setInputValueForParamAndTriggerChange('forecast-length', '{fhr}');
                                matsParamUtils.setInputValueForParamAndTriggerChange('curve-dates', '{sfc_dates}');
                            }}''')
                            await pg.wait_for_timeout(150)
                            await pg.click("#add")
                            await pg.wait_for_timeout(150)
                        await pg.click("#plotMatched")
                        await pg.wait_for_timeout(8000)
                        await export_styled_chart(pg, out_bias, color_map, models=valid_models)
                    except Exception as e:
                        print(f"  [WARNING] Failed to fetch Visibility {fhr}-h TS Bias ({tag}, {reg}): {e}. Skipping.")
                    finally:
                        await b.close()
                else:
                    print(f"  [CACHE HIT] Using cached: {fn_bias}")


def build_presentation(config, mats_dir, output_pptx):
    """
    Builds the PowerPoint presentation deck matching NOAA / retv4 standards.
    Collects and returns structured slide metadata `slides_data` for exact HTML synchronization.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN
        from pptx.dml.color import RGBColor
    except ImportError:
        print("Error: python-pptx is not installed. Run `pip install python-pptx`")
        return []

    print(f"\n>>> Compiling PowerPoint presentation: {output_pptx}...")
    prs = Presentation()
    prs.slide_width = 9144000
    prs.slide_height = 5143500
    blank_layout = prs.slide_layouts[6]
    
    models = config.get("models", [])
    title_text = config.get("title", "Numerical Weather Prediction Verification Report")
    sfc_dates = config.get("dates", {}).get("surface", "")
    ua_dates = config.get("dates", {}).get("upperair", "")
    cref_dir = config.get("cref_dir") or os.path.join(os.path.dirname(os.path.abspath(output_pptx)), "cref")
    extracted_dir = config.get("extracted_ref_dir") or os.path.join(os.path.dirname(os.path.abspath(output_pptx)), "extracted_retv4")
    
    slides_data = []

    def add_header(slide, title):
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(9.0), Inches(0.45))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(19)
        p.font.bold = True
        p.font.color.rgb = RGBColor(16, 44, 87)
        p.alignment = PP_ALIGN.CENTER

    def find_img(fn):
        if not fn: return None
        p = os.path.join(mats_dir, fn)
        if os.path.exists(p): return p
        if "ft." in fn:
            p_alt = os.path.join(mats_dir, fn.replace("ft.", "."))
            if os.path.exists(p_alt): return p_alt
        if not "ft." in fn and any(k in fn for k in ["500.", "1000.", "3000."]):
            p_alt = os.path.join(mats_dir, fn.replace("500.", "500ft.").replace("1000.", "1000ft.").replace("3000.", "3000ft."))
            if os.path.exists(p_alt): return p_alt
        return None

    def add_legend_box(slide, is_bias=False, metric="RMSE"):
        box = slide.shapes.add_textbox(7313400, 2958262, 1650000, 1600000)
        tf = box.text_frame
        tf.word_wrap = True
        lines = [("Experiments:", True, Pt(11))]
        for m in models:
            lines.append((f"• {m['name']}", False, Pt(10)))
        if is_bias:
            lines.append(("Metric:", True, Pt(11)))
            lines.append(("• Dash: BIAS", False, Pt(10)))
        elif metric == "COMBINED":
            lines.append(("Styles:", True, Pt(11)))
            lines.append(("• Solid: RMSE", False, Pt(10)))
            lines.append(("• Dash: BIAS", False, Pt(10)))
        else:
            lines.append(("Metric:", True, Pt(11)))
            lines.append((f"• Solid: {metric}", False, Pt(10)))
        for t, b, s in lines:
            p = tf.add_paragraph()
            p.text = t
            p.font.bold = b
            p.font.size = s

    def add_3panel_slide(title, f1, f2, f3, l1="t2m", l2="td2m", l3="wind10m", is_bias=False, metric="RMSE", region="all"):
        p1 = find_img(f1)
        p2 = find_img(f2)
        p3 = find_img(f3)
        if not (p1 or p2 or p3):
            return False
            
        slide = prs.slides.add_slide(blank_layout)
        add_header(slide, title)
        
        if p1:
            slide.shapes.add_picture(p1, 216610, 705325, 4244023, 2127095)
            lbl = slide.shapes.add_textbox(1886891, 355817, 1200000, 350000)
            lbl.text_frame.text = l1
            lbl.text_frame.paragraphs[0].font.bold = True
            
        if p2:
            slide.shapes.add_picture(p2, 4506775, 714750, 4529823, 2127100)
            lbl = slide.shapes.add_textbox(5388838, 378849, 1200000, 350000)
            lbl.text_frame.text = l2
            lbl.text_frame.paragraphs[0].font.bold = True
            
        if p3:
            slide.shapes.add_picture(p3, 2365625, 2894825, 4529823, 2243827)
            lbl = slide.shapes.add_textbox(1468986, 3620929, 1200000, 350000)
            lbl.text_frame.text = l3
            lbl.text_frame.paragraphs[0].font.bold = True
            
        add_legend_box(slide, is_bias=is_bias, metric=metric)
        slides_data.append({
            "type": "3panel",
            "title": title,
            "img1": p1, "lbl1": l1,
            "img2": p2, "lbl2": l2,
            "img3": p3, "lbl3": l3,
            "is_bias": is_bias,
            "metric": metric,
            "models": models,
            "region": region
        })
        return True

    def add_integrated_ts_slide(title, col1_name, col2_name, row_specs, region="all"):
        active_rows = []
        for lbl, f1, f2 in row_specs:
            p1 = find_img(f1)
            p2 = find_img(f2)
            if p1 or p2:
                active_rows.append({"label": lbl, "img1": p1, "img2": p2})
                
        if not active_rows:
            return False

        slide = prs.slides.add_slide(blank_layout)
        add_header(slide, title)
        
        # Dimensions (in EMU): Slide W: 9144000, Slide H: 5143500
        x_lbl = int(0.20 * 914400)
        w_lbl = int(1.15 * 914400)
        
        x_col1 = int(1.40 * 914400)
        w_col = int(3.35 * 914400)
        
        x_col2 = int(4.85 * 914400)
        
        x_leg = int(8.30 * 914400)
        w_leg = int(1.65 * 914400)
        
        y_hdr = int(0.55 * 914400)
        h_hdr = int(0.28 * 914400)
        
        # Column headers
        t1 = slide.shapes.add_textbox(x_col1, y_hdr, w_col, h_hdr)
        t1.text_frame.text = col1_name
        t1.text_frame.paragraphs[0].font.bold = True
        t1.text_frame.paragraphs[0].font.size = Pt(13)
        t1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        t2 = slide.shapes.add_textbox(x_col2, y_hdr, w_col, h_hdr)
        t2.text_frame.text = col2_name
        t2.text_frame.paragraphs[0].font.bold = True
        t2.text_frame.paragraphs[0].font.size = Pt(13)
        t2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 3 Rows positioned with zero overflow:
        h_pic = int(1.44 * 914400)
        gap_y = int(0.06 * 914400)
        y_start = int(0.88 * 914400)
        
        for idx, r in enumerate(active_rows):
            y_pos = y_start + idx * (h_pic + gap_y)
            
            lbl_box = slide.shapes.add_textbox(x_lbl, y_pos + int(h_pic/2) - int(0.18*914400), w_lbl, int(0.36*914400))
            lbl_box.text_frame.text = r["label"]
            lbl_box.text_frame.paragraphs[0].font.bold = True
            lbl_box.text_frame.paragraphs[0].font.size = Pt(11)
            lbl_box.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            
            if r["img1"]:
                slide.shapes.add_picture(r["img1"], x_col1, y_pos, w_col, h_pic)
            if r["img2"]:
                slide.shapes.add_picture(r["img2"], x_col2, y_pos, w_col, h_pic)
                
        # Legend Box on right
        leg = slide.shapes.add_textbox(x_leg, int(1.10 * 914400), w_leg, int(3.80 * 914400))
        tf = leg.text_frame
        tf.word_wrap = True
        lines = [("Experiments:", True, Pt(11))]
        for m in models:
            lines.append((f"• {m['name']}", False, Pt(10)))
        lines.append(("Columns:", True, Pt(11)))
        lines.append((f"• Left: {col1_name}", False, Pt(10)))
        lines.append((f"• Right: {col2_name}", False, Pt(10)))
        lines.append(("Line Styles:", True, Pt(11)))
        lines.append((f"• Solid: {col1_name}", False, Pt(10)))
        lines.append((f"• Dash: {col2_name}", False, Pt(10)))
        
        for t, b, s in lines:
            p = tf.add_paragraph()
            p.text = t
            p.font.bold = b
            p.font.size = s
            
        slides_data.append({
            "type": "integrated_ts",
            "title": title,
            "col1_name": col1_name,
            "col2_name": col2_name,
            "rows": active_rows,
            "models": models,
            "region": region
        })
        return True

    def add_profile_slide(title, f_t, f_w, f_rh, region="all"):
        p_t = find_img(f_t)
        p_w = find_img(f_w)
        p_rh = find_img(f_rh)
        if not (p_t or p_w or p_rh):
            return False
            
        sl = prs.slides.add_slide(blank_layout)
        add_header(sl, title)
        
        if p_rh:
            # 3-plot case (RAOBs: Temp, Wind, RHobT + Legend) -> 2x2 grid
            c1_x = Inches(2.35)
            c2_x = Inches(5.15)
            p_size = Inches(2.00)
            
            # Row 1
            if p_t:
                lbl1 = sl.shapes.add_textbox(c1_x, Inches(0.56), p_size, Inches(0.24))
                lbl1.text_frame.text = "Temperature (t)"
                lbl1.text_frame.paragraphs[0].font.bold = True
                lbl1.text_frame.paragraphs[0].font.size = Pt(11)
                lbl1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                sl.shapes.add_picture(p_t, c1_x, Inches(0.80), p_size, p_size)
            if p_w:
                lbl2 = sl.shapes.add_textbox(c2_x, Inches(0.56), p_size, Inches(0.24))
                lbl2.text_frame.text = "Wind"
                lbl2.text_frame.paragraphs[0].font.bold = True
                lbl2.text_frame.paragraphs[0].font.size = Pt(11)
                lbl2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                sl.shapes.add_picture(p_w, c2_x, Inches(0.80), p_size, p_size)
                
            # Row 2
            lbl3 = sl.shapes.add_textbox(c1_x, Inches(2.95), p_size, Inches(0.24))
            lbl3.text_frame.text = "Relative Humidity (RHobT)"
            lbl3.text_frame.paragraphs[0].font.bold = True
            lbl3.text_frame.paragraphs[0].font.size = Pt(11)
            lbl3.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            sl.shapes.add_picture(p_rh, c1_x, Inches(3.20), p_size, p_size)
            
            # Legend at Row 2, Col 2
            leg = sl.shapes.add_textbox(c2_x, Inches(3.20), Inches(2.45), Inches(1.95))
        else:
            # 2-plot case (AMDAR: Temp, Wind + Legend) -> 3-col side-by-side
            c1_x = Inches(0.50)
            c2_x = Inches(4.20)
            c3_x = Inches(7.90)
            p_size = Inches(3.45)
            
            if p_t:
                lbl1 = sl.shapes.add_textbox(c1_x, Inches(0.58), p_size, Inches(0.28))
                lbl1.text_frame.text = "Temperature (t)"
                lbl1.text_frame.paragraphs[0].font.bold = True
                lbl1.text_frame.paragraphs[0].font.size = Pt(12)
                lbl1.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                sl.shapes.add_picture(p_t, c1_x, Inches(0.88), p_size, p_size)
            if p_w:
                lbl2 = sl.shapes.add_textbox(c2_x, Inches(0.58), p_size, Inches(0.28))
                lbl2.text_frame.text = "Wind"
                lbl2.text_frame.paragraphs[0].font.bold = True
                lbl2.text_frame.paragraphs[0].font.size = Pt(12)
                lbl2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                sl.shapes.add_picture(p_w, c2_x, Inches(0.88), p_size, p_size)
                
            # Legend at Col 3
            leg = sl.shapes.add_textbox(c3_x, Inches(0.88), Inches(1.65), p_size)
            
        tf = leg.text_frame
        tf.word_wrap = True
        lines = [("Experiments:", True, Pt(11))]
        for m in models: lines.append((f"• {m['name']}", False, Pt(10)))
        lines.append(("Line Styles:", True, Pt(11)))
        lines.append(("• Solid: RMSE", False, Pt(10)))
        lines.append(("• Dash: BIAS", False, Pt(10)))
        for t, b, s in lines:
            p = tf.add_paragraph()
            p.text = t
            p.font.bold = b
            p.font.size = s
            
        slides_data.append({
            "type": "profile",
            "title": title,
            "img_t": p_t, "img_w": p_w, "img_rh": p_rh,
            "models": models,
            "region": region
        })
        return True

    def add_section_divider_slide(region_id):
        slide = prs.slides.add_slide(blank_layout)
        # Background full-bleed dark rectangle
        shape = slide.shapes.add_shape(
            1, # MSO_SHAPE.RECTANGLE
            0, 0, prs.slide_width, prs.slide_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(16, 44, 87)
        shape.line.fill.background()

        reg_full_name = REGION_NAMES.get(region_id, region_id.capitalize())
        
        # Pill Header
        tb_pill = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(8.0), Inches(0.45))
        p_pill = tb_pill.text_frame.paragraphs[0]
        p_pill.text = "REGIONAL VERIFICATION CHAPTER"
        p_pill.font.size = Pt(13)
        p_pill.font.bold = True
        p_pill.font.color.rgb = RGBColor(56, 189, 248) # Sky blue

        # Main Title
        tb_title = slide.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(8.0), Inches(1.3))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = f"Sub-Domain Evaluation:\n{reg_full_name}"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

        # Subtitle / Details
        tb_sub = slide.shapes.add_textbox(Inches(1.0), Inches(3.25), Inches(8.0), Inches(0.8))
        p_sub = tb_sub.text_frame.paragraphs[0]
        p_sub.text = f"Comparative performance analysis across {reg_full_name} for Surface, Upper Air, Radar, and Aviation METARs."
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = RGBColor(203, 213, 225)

        slides_data.append({
            "type": "section_header",
            "region": region_id,
            "title": f"Sub-Domain Evaluation: {reg_full_name}",
            "subtitle": f"Comparative performance analysis across {reg_full_name}."
        })

    def add_regional_verification_deck(reg):
        reg_tag = "" if reg == "all" else f"_{reg}"
        reg_suffix = "" if reg == "all" else f" ({REGION_NAMES.get(reg, reg).split(' (')[0]})"

        # 1. Surface Dieoff
        add_3panel_slide(
            f"Surface Dieoff{reg_suffix}",
            f"surface_dieoff_t2m{reg_tag}.png",
            f"surface_dieoff_td2m{reg_tag}.png",
            f"surface_dieoff_wind10m{reg_tag}.png",
            metric="COMBINED",
            region=reg
        )

        # 2. Surface Time Series
        for ts_title, f_lead in [
            (f"Surface 1-h fcst time series{reg_suffix}", "1h"),
            (f"Surface 3-h fcst time series{reg_suffix}", "3h"),
            (f"Surface 12-h fcst time series{reg_suffix}", "12h"),
            (f"Surface 24-h fcst time series{reg_suffix}", "24h")
        ]:
            add_integrated_ts_slide(
                ts_title, "RMSE", "BIAS",
                [
                    ("t2m", f"sfc_ts_{f_lead}_rmse_t2m{reg_tag}.png", f"sfc_ts_{f_lead}_bias_t2m{reg_tag}.png"),
                    ("td2m", f"sfc_ts_{f_lead}_rmse_td2m{reg_tag}.png", f"sfc_ts_{f_lead}_bias_td2m{reg_tag}.png"),
                    ("Wind 10m", f"sfc_ts_{f_lead}_rmse_wind10m{reg_tag}.png", f"sfc_ts_{f_lead}_bias_wind10m{reg_tag}.png")
                ],
                region=reg
            )

        # 3. Upper Air RAOBs Dieoff
        add_3panel_slide(
            f"Upper air against RAOBs dieoff{reg_suffix}",
            f"ua_dieoff_temp{reg_tag}.png",
            f"ua_dieoff_rh{reg_tag}.png",
            f"ua_dieoff_wind{reg_tag}.png",
            l1="Temperature", l2="RHobT", l3="Wind", metric="COMBINED",
            region=reg
        )

        # AMDAR Dieoff (AMDAR is only available for full domain)
        if reg == "all":
            add_3panel_slide("Upper air against AMDAR dieoff", "amdar_dieoff_temp.png", "amdar_dieoff_rh.png", "amdar_dieoff_wind.png", l1="Temperature", l2="RHobT", l3="Wind", metric="COMBINED", region=reg)

        # 4. Upper Air RAOBs Profiles (1h, 3h, 6h, 12h)
        add_profile_slide(f"Upper air 1-h fcst (RAOBs Profile){reg_suffix}", f"ua_prof_1h_temp{reg_tag}.png", f"ua_prof_1h_wind{reg_tag}.png", f"ua_prof_1h_rh{reg_tag}.png", region=reg)
        add_profile_slide(f"Upper air 3-h fcst (RAOBs Profile){reg_suffix}", f"ua_prof_3h_temp{reg_tag}.png", f"ua_prof_3h_wind{reg_tag}.png", f"ua_prof_3h_rh{reg_tag}.png", region=reg)
        add_profile_slide(f"Upper air 6-h fcst (RAOBs Profile){reg_suffix}", f"ua_prof_6h_temp{reg_tag}.png", f"ua_prof_6h_wind{reg_tag}.png", f"ua_prof_6h_rh{reg_tag}.png", region=reg)
        add_profile_slide(f"Upper air 12-h fcst (RAOBs Profile){reg_suffix}", f"ua_prof_12h_temp{reg_tag}.png", f"ua_prof_12h_wind{reg_tag}.png", f"ua_prof_12h_rh{reg_tag}.png", region=reg)

        # AMDAR Profiles (full domain only)
        if reg == "all":
            add_profile_slide("Upper air 1-h fcst (AMDAR Profile)", "amdar_prof_1h_temp.png", "amdar_prof_1h_wind.png", "amdar_prof_1h_rh.png", region=reg)
            add_profile_slide("Upper air 3-h fcst (AMDAR Profile)", "amdar_prof_3h_temp.png", "amdar_prof_3h_wind.png", "amdar_prof_3h_rh.png", region=reg)
            add_profile_slide("Upper air 6-h fcst (AMDAR Profile)", "amdar_prof_6h_temp.png", "amdar_prof_6h_wind.png", "amdar_prof_6h_rh.png", region=reg)
            add_profile_slide("Upper air 12-h fcst (AMDAR Profile)", "amdar_prof_12h_temp.png", "amdar_prof_12h_wind.png", "amdar_prof_12h_rh.png", region=reg)

        # 5. Radar Reflectivity Verification
        add_3panel_slide(f"CREF Dieoff (CSI ){reg_suffix}", f"ref_dieoff_csi_15dbz{reg_tag}.png", f"ref_dieoff_csi_25dbz{reg_tag}.png", f"ref_dieoff_csi_35dbz{reg_tag}.png", l1="15 dBZ", l2="25 dBZ", l3="35 dBZ", metric="CSI", region=reg)
        add_3panel_slide(f"CREF Dieoff (BIAS ){reg_suffix}", f"ref_dieoff_bias_15dbz{reg_tag}.png", f"ref_dieoff_bias_25dbz{reg_tag}.png", f"ref_dieoff_bias_35dbz{reg_tag}.png", l1="15 dBZ", l2="25 dBZ", l3="35 dBZ", is_bias=True, region=reg)
        for ts_title, f_lead in [(f"CREF 1-h fcst time series{reg_suffix}", "1h"), (f"CREF 3-h fcst time series{reg_suffix}", "3h")]:
            add_integrated_ts_slide(
                ts_title, "CSI", "BIAS",
                [
                    ("15 dBZ", f"ref_ts_{f_lead}_csi_15dbz{reg_tag}.png", f"ref_ts_{f_lead}_bias_15dbz{reg_tag}.png"),
                    ("25 dBZ", f"ref_ts_{f_lead}_csi_25dbz{reg_tag}.png", f"ref_ts_{f_lead}_bias_25dbz{reg_tag}.png"),
                    ("35 dBZ", f"ref_ts_{f_lead}_csi_35dbz{reg_tag}.png", f"ref_ts_{f_lead}_bias_35dbz{reg_tag}.png")
                ],
                region=reg
            )

        # 6. Ceiling Verification
        add_3panel_slide(f"Ceiling Dieoff (CSI){reg_suffix}", f"ceil_dieoff_csi_500ft{reg_tag}.png", f"ceil_dieoff_csi_1000ft{reg_tag}.png", f"ceil_dieoff_csi_3000ft{reg_tag}.png", l1="< 500 ft", l2="< 1000 ft", l3="< 3000 ft", metric="CSI", region=reg)
        add_3panel_slide(f"Ceiling Dieoff (BIAS){reg_suffix}", f"ceil_dieoff_bias_500ft{reg_tag}.png", f"ceil_dieoff_bias_1000ft{reg_tag}.png", f"ceil_dieoff_bias_3000ft{reg_tag}.png", l1="< 500 ft", l2="< 1000 ft", l3="< 3000 ft", is_bias=True, region=reg)
        for ts_title, f_lead in [(f"Ceiling 1-h fcst time series{reg_suffix}", "1h"), (f"Ceiling 3-h fcst time series{reg_suffix}", "3h")]:
            add_integrated_ts_slide(
                ts_title, "CSI", "BIAS",
                [
                    ("< 500 ft", f"ceil_ts_{f_lead}_csi_500ft{reg_tag}.png", f"ceil_ts_{f_lead}_bias_500ft{reg_tag}.png"),
                    ("< 1000 ft", f"ceil_ts_{f_lead}_csi_1000ft{reg_tag}.png", f"ceil_ts_{f_lead}_bias_1000ft{reg_tag}.png"),
                    ("< 3000 ft", f"ceil_ts_{f_lead}_csi_3000ft{reg_tag}.png", f"ceil_ts_{f_lead}_bias_3000ft{reg_tag}.png")
                ],
                region=reg
            )

        # 7. Visibility Verification
        add_3panel_slide(f"Visibility Dieoff (CSI){reg_suffix}", f"vis_dieoff_csi_1mi{reg_tag}.png", f"vis_dieoff_csi_3mi{reg_tag}.png", f"vis_dieoff_csi_5mi{reg_tag}.png", l1="< 1 mi", l2="< 3 mi", l3="< 5 mi", metric="CSI", region=reg)
        add_3panel_slide(f"Visibility Dieoff (BIAS){reg_suffix}", f"vis_dieoff_bias_1mi{reg_tag}.png", f"vis_dieoff_bias_3mi{reg_tag}.png", f"vis_dieoff_bias_5mi{reg_tag}.png", l1="< 1 mi", l2="< 3 mi", l3="< 5 mi", is_bias=True, region=reg)
        for ts_title, f_lead in [(f"Visibility 1-h fcst time series{reg_suffix}", "1h"), (f"Visibility 3-h fcst time series{reg_suffix}", "3h")]:
            add_integrated_ts_slide(
                ts_title, "CSI", "BIAS",
                [
                    ("< 1 mi", f"vis_ts_{f_lead}_csi_1mi{reg_tag}.png", f"vis_ts_{f_lead}_bias_1mi{reg_tag}.png"),
                    ("< 3 mi", f"vis_ts_{f_lead}_csi_3mi{reg_tag}.png", f"vis_ts_{f_lead}_bias_3mi{reg_tag}.png"),
                    ("< 5 mi", f"vis_ts_{f_lead}_csi_5mi{reg_tag}.png", f"vis_ts_{f_lead}_bias_5mi{reg_tag}.png")
                ],
                region=reg
            )

    regions = config.get("regions", ["all"])

    # Slide 1: Title
    s1 = prs.slides.add_slide(blank_layout)
    tb = s1.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = RGBColor(16, 44, 87)
    
    b_box = s1.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(8.8), Inches(3.5))
    tf = b_box.text_frame
    tf.word_wrap = True
    bullets = [
        ("Experiments Evaluated:", True),
    ]
    for m in models:
        bullets.append((f"{m['name']} ({m['id']})", False))
    reg_names_str = ", ".join([REGION_NAMES.get(r, r).split(" (")[0] for r in regions])
    bullets.append((f"Evaluated Domains: {reg_names_str}", True))
    bullets.extend([
        ("Verification Sources:", True),
        ("METAR surface station observations (Plot Matched)", False),
        ("Traditional Rawinsonde (RAOBs) soundings from 1000 hPa to 100 hPa", False),
        ("Cloud Ceiling & Surface Visibility (METAR thresholds)", False),
    ])
    cases = get_radar_cases(config)
    if any(find_img(f) for f in ["ref_dieoff_csi_15dbz.png", "ref_ts_1h_csi_15dbz.png"]) or cases:
        bullets.append(("Composite Reflectivity (MRMS Radar at 15, 25, 35 dBZ)", False))
    if any(find_img(f) for f in ["amdar_dieoff_temp.png", "amdar_dieoff_wind.png"]):
        bullets.append(("Aircraft observations (AMDAR traditional)", False))
    for text, bold in bullets:
        p = tf.add_paragraph()
        if bold:
            p.text = text
            p.font.size = Pt(13)
            p.font.bold = True
        else:
            p.text = f"  • {text}"
            p.font.size = Pt(13)
            p.font.bold = False

    slides_data.append({
        "type": "title",
        "title": title_text,
        "models": models,
        "bullets": bullets,
        "dates": {"surface": sfc_dates, "upperair": ua_dates},
        "regions": regions,
        "region": "all"
    })

    # Baseline: Full Domain Verification Deck
    add_regional_verification_deck("all")

    # Additional Regional Sections (Scheme A: Independent Section Divider + Sub-Domain Slides)
    for reg in regions:
        if reg != "all":
            add_section_divider_slide(reg)
            add_regional_verification_deck(reg)

    # Synoptic Radar Reflectivity Comparison Cases
    # Dynamically generated from configured radar_cases or DEFAULT_RADAR_CASES
    if config.get("skip_cref"):
        cases = []

    w = int(2.85 * 914400)
    h = int(2.15 * 914400)
    y_top = int(0.55 * 914400)
    x1, x2, x3 = int(0.4 * 914400), int(3.4 * 914400), int(6.4 * 914400)
    y_bot = int(2.85 * 914400)

    for orig_idx, cdate, cycle, fhr, case_title in cases:
        # 1. MRMS observation
        mrms_fn = f"mrms_{cdate}_{cycle}z_f{fhr:02d}.png"
        p_mrms_cref = os.path.join(cref_dir, mrms_fn)
        if os.path.exists(p_mrms_cref) and os.path.getsize(p_mrms_cref) > 10000:
            p_mrms = p_mrms_cref
        elif os.path.exists(extracted_dir):
            legacy_p = os.path.join(extracted_dir, f"slide_{orig_idx}_img_3.png")
            p_mrms = legacy_p if (os.path.exists(legacy_p) and os.path.getsize(legacy_p) > 10000) else ""
        else:
            p_mrms = ""

        # 2. Build cards list: Panel 1 is always MRMS (Observation)
        cards = [("MRMS (Observation)", p_mrms if (p_mrms and os.path.exists(p_mrms)) else None)]

        # 3. Add each evaluated model from models
        model_names_lower = []
        for m in models:
            m_name = m['name']
            model_names_lower.append(m_name.lower())
            m_fn = f"{m_name}_{cdate}_{cycle}z_f{fhr:02d}.png"
            m_path = os.path.join(cref_dir, m_fn)
            found_img = None
            if os.path.exists(m_path) and os.path.getsize(m_path) > 1000:
                found_img = m_path
            elif os.path.exists(extracted_dir):
                if m_name == "retv3":
                    cand = os.path.join(extracted_dir, f"slide_{orig_idx}_img_1.png")
                    if os.path.exists(cand) and os.path.getsize(cand) > 1000:
                        found_img = cand
                elif m_name == "retv4":
                    cand = os.path.join(extracted_dir, f"slide_{orig_idx}_img_2.png")
                    if os.path.exists(cand) and os.path.getsize(cand) > 1000:
                        found_img = cand
            cards.append((m_name, found_img))

        # If extracted_dir has retv3/retv4 but they were not in models list, append as baselines if images exist
        if "retv3" not in model_names_lower and os.path.exists(extracted_dir):
            cand = os.path.join(extracted_dir, f"slide_{orig_idx}_img_1.png")
            if os.path.exists(cand) and os.path.getsize(cand) > 1000:
                cards.append(("retv3 (Baseline 3)", cand))
        if "retv4" not in model_names_lower and os.path.exists(extracted_dir):
            cand = os.path.join(extracted_dir, f"slide_{orig_idx}_img_2.png")
            if os.path.exists(cand) and os.path.getsize(cand) > 1000:
                cards.append(("retv4 (Baseline 4)", cand))

        # Check if any image exists in this case
        has_any_img = any(img_p for _, img_p in cards)
        if not has_any_img:
            continue

        s_radar = prs.slides.add_slide(blank_layout)
        t_box = s_radar.shapes.add_textbox(Inches(0.4), Inches(0.1), Inches(8.5), Inches(0.4))
        p = t_box.text_frame.paragraphs[0]
        p.text = f"Composite Reflectivity (dBZ) Comparison: {case_title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(16, 44, 87)

        # Layout up to 6 cards in a 3x2 grid
        grid_coords = [
            (x1, y_top), (x2, y_top), (x3, y_top),
            (x1, y_bot), (x2, y_bot), (x3, y_bot)
        ]

        for idx, (label_text, img_p) in enumerate(cards[:6]):
            xc, yc = grid_coords[idx]

            # Label box
            lbl = s_radar.shapes.add_textbox(xc, yc - int(0.25 * 914400), w, int(0.25 * 914400))
            lbl.text_frame.text = label_text
            lbl.text_frame.paragraphs[0].font.bold = True
            lbl.text_frame.paragraphs[0].font.size = Pt(11)

            if img_p and os.path.exists(img_p):
                s_radar.shapes.add_picture(img_p, xc, yc, w, h)
            else:
                placeholder = s_radar.shapes.add_textbox(xc, yc, w, h)
                placeholder.fill.solid()
                placeholder.fill.fore_color.rgb = RGBColor(241, 245, 249)
                placeholder.line.color.rgb = RGBColor(203, 213, 225)
                tf_ph = placeholder.text_frame
                tf_ph.word_wrap = True
                p_ph = tf_ph.paragraphs[0]
                p_ph.text = "\n\nNo Image Available"
                p_ph.alignment = PP_ALIGN.CENTER
                p_ph.font.size = Pt(11)
                p_ph.font.color.rgb = RGBColor(148, 163, 184)

        slides_data.append({
            "type": "radar_case",
            "title": f"Composite Reflectivity (dBZ) Comparison: {case_title}",
            "case_title": case_title,
            "cards": cards,
            "region": "cases"
        })

    atomic_save_pptx(prs, output_pptx)
    print(f"\n[DONE] Successfully created PowerPoint presentation: {output_pptx}")
    return slides_data


def build_html_report(config, mats_dir, output_html, slides_data):
    """
    Renders an identical, interactive HTML report with exact 1-to-1 slide alignment to the PPT presentation.
    Fixes legend dashed lines with crisp SVG vectors, and supports slideshow & scroll presentation modes.
    """
    print(f"\n>>> Compiling HTML Presentation Report: {output_html}...")
    html_dir = os.path.dirname(os.path.abspath(output_html))
    os.makedirs(html_dir, exist_ok=True)
    
    def rel_img(path):
        if not path:
            return ""
        try:
            return os.path.relpath(os.path.abspath(path), html_dir)
        except Exception:
            return path

    title_text = config.get("title", "Numerical Weather Prediction Verification Report")
    
    slides_html = []
    nav_options = []

    # SVG helper for crisp solid & dashed legend samples
    svg_solid = '<svg width="24" height="6" style="vertical-align: middle;"><line x1="0" y1="3" x2="24" y2="3" stroke="#334155" stroke-width="3" /></svg>'
    svg_dash = '<svg width="24" height="6" style="vertical-align: middle;"><line x1="0" y1="3" x2="24" y2="3" stroke="#334155" stroke-width="3" stroke-dasharray="5,3" /></svg>'

    for idx, slide in enumerate(slides_data, start=1):
        stype = slide.get("type")
        stitle = slide.get("title", f"Slide {idx}")
        safe_title = html.escape(stitle)
        nav_options.append(f'<option value="{idx}">{idx}. {safe_title}</option>')
        
        slide_content = ""
        
        if stype == "title":
            bullets_html = []
            for b_text, is_bold in slide.get("bullets", []):
                clean_text = b_text.strip().lstrip("•").strip()
                s_text = html.escape(clean_text)
                if is_bold:
                    bullets_html.append(f'<li class="bullet-header"><strong>{s_text}</strong></li>')
                else:
                    bullets_html.append(f'<li class="bullet-item">{s_text}</li>')
                    
            model_chips = "".join([
                f'<span class="model-badge" style="border-left: 5px solid {m.get("color", "#2563eb")}">'
                f'<strong>{html.escape(m["name"])}</strong> <small>({html.escape(m["id"])})</small></span>'
                for m in slide.get("models", [])
            ])
            
            sfc_d = html.escape(slide.get("dates", {}).get("surface", "N/A"))
            ua_d = html.escape(slide.get("dates", {}).get("upperair", "N/A"))
            
            slide_content = f'''
            <div class="title-slide-container">
              <div class="title-hero">
                <span class="report-tag">NOAA GSL · Verification Executive Briefing</span>
                <h1 class="main-title">{safe_title}</h1>
              </div>
              <div class="title-body">
                <div class="models-panel">
                  <h3>Evaluated Experiments:</h3>
                  <div class="models-grid">{model_chips}</div>
                  <div class="date-ranges">
                    <div><strong>Surface Period:</strong> {sfc_d}</div>
                    <div><strong>Upper-Air Period:</strong> {ua_d}</div>
                  </div>
                </div>
                <div class="sources-panel">
                  <h3>Verification Framework:</h3>
                  <ul class="bullet-list">{ "".join(bullets_html) }</ul>
                </div>
              </div>
            </div>
            '''
            
        elif stype == "3panel":
            img1 = rel_img(slide.get("img1"))
            img2 = rel_img(slide.get("img2"))
            img3 = rel_img(slide.get("img3"))
            lbl1 = html.escape(slide.get("lbl1", ""))
            lbl2 = html.escape(slide.get("lbl2", ""))
            lbl3 = html.escape(slide.get("lbl3", ""))
            metric = html.escape(slide.get("metric", "RMSE"))
            is_bias = slide.get("is_bias", False)
            
            exp_lines = "".join([
                f'<div class="legend-row"><span class="color-dot" style="background:{m.get("color", "#64748b")}"></span><span>{html.escape(m["name"])}</span></div>'
                for m in slide.get("models", [])
            ])
            
            if is_bias:
                style_legend = f'<div class="legend-subtitle">Metric:</div><div class="legend-row">{svg_dash}<span>Dash: BIAS</span></div>'
            elif metric == "COMBINED":
                style_legend = f'<div class="legend-subtitle">Line Styles:</div><div class="legend-row">{svg_solid}<span>Solid: RMSE</span></div><div class="legend-row">{svg_dash}<span>Dash: BIAS</span></div>'
            else:
                style_legend = f'<div class="legend-subtitle">Metric:</div><div class="legend-row">{svg_solid}<span>Solid: {metric}</span></div>'

            slide_content = f'''
            <div class="slide-header"><h2>{safe_title}</h2></div>
            <div class="grid-3panel">
              <div class="panel-box p1">
                <div class="panel-label">{lbl1}</div>
                {f'<img src="{img1}" alt="{lbl1}" onclick="openLightbox(this)">' if img1 else '<div class="no-img">No Image</div>'}
              </div>
              <div class="panel-box p2">
                <div class="panel-label">{lbl2}</div>
                {f'<img src="{img2}" alt="{lbl2}" onclick="openLightbox(this)">' if img2 else '<div class="no-img">No Image</div>'}
              </div>
              <div class="panel-box p3">
                <div class="panel-label">{lbl3}</div>
                {f'<img src="{img3}" alt="{lbl3}" onclick="openLightbox(this)">' if img3 else '<div class="no-img">No Image</div>'}
              </div>
              <div class="panel-box legend-box">
                <div class="legend-title">Experiments:</div>
                {exp_lines}
                {style_legend}
              </div>
            </div>
            '''
            
        elif stype in ["integrated_ts", "sfc24"]:
            col1_name = html.escape(slide.get("col1_name", "RMSE"))
            col2_name = html.escape(slide.get("col2_name", "BIAS"))
            row_html = []
            for r in slide.get("rows", []):
                rlbl = html.escape(r.get("label", ""))
                r_img1 = rel_img(r.get("img1") or r.get("img_rmse"))
                r_img2 = rel_img(r.get("img2") or r.get("img_bias"))
                row_html.append(f'''
                <div class="ts-row">
                  <div class="ts-row-label">{rlbl}</div>
                  <div class="ts-img-box">{f'<img src="{r_img1}" alt="{rlbl} {col1_name}" onclick="openLightbox(this)">' if r_img1 else '<div class="no-img">No Image</div>'}</div>
                  <div class="ts-img-box">{f'<img src="{r_img2}" alt="{rlbl} {col2_name}" onclick="openLightbox(this)">' if r_img2 else '<div class="no-img">No Image</div>'}</div>
                </div>
                ''')

            exp_lines = "".join([
                f'<div class="legend-row"><span class="color-dot" style="background:{m.get("color", "#64748b")}"></span><span>{html.escape(m["name"])}</span></div>'
                for m in slide.get("models", [])
            ])

            slide_content = f'''
            <div class="slide-header"><h2>{safe_title}</h2></div>
            <div class="ts-integrated-container">
              <div class="ts-main-grid">
                <div class="ts-header-row">
                  <div class="ts-label-space"></div>
                  <div class="ts-col-header">{col1_name}</div>
                  <div class="ts-col-header">{col2_name}</div>
                </div>
                {"".join(row_html)}
              </div>
              <div class="ts-legend-sidebar">
                <div class="legend-title">Experiments:</div>
                {exp_lines}
                <div class="legend-subtitle">Columns:</div>
                <div class="legend-row"><span>• Left: {col1_name}</span></div>
                <div class="legend-row"><span>• Right: {col2_name}</span></div>
                <div class="legend-subtitle">Line Styles:</div>
                <div class="legend-row">{svg_solid}<span>Solid: {col1_name}</span></div>
                <div class="legend-row">{svg_dash}<span>Dash: {col2_name}</span></div>
              </div>
            </div>
            '''
            
        elif stype == "profile":
            img_t = rel_img(slide.get("img_t"))
            img_w = rel_img(slide.get("img_w"))
            img_rh = rel_img(slide.get("img_rh"))
            exp_lines = "".join([
                f'<div class="legend-row"><span class="color-dot" style="background:{m.get("color", "#64748b")}"></span><span>{html.escape(m["name"])}</span></div>'
                for m in slide.get("models", [])
            ])
            if img_rh:
                slide_content = f'''
                <div class="slide-header"><h2>{safe_title}</h2></div>
                <div class="grid-profile">
                  <div class="profile-box p-temp">
                    <div class="profile-label">Temperature (t)</div>
                    {f'<img src="{img_t}" alt="Temperature Profile" onclick="openLightbox(this)">' if img_t else '<div class="no-img">No Image</div>'}
                  </div>
                  <div class="profile-box p-wind">
                    <div class="profile-label">Wind</div>
                    {f'<img src="{img_w}" alt="Wind Profile" onclick="openLightbox(this)">' if img_w else '<div class="no-img">No Image</div>'}
                  </div>
                  <div class="profile-box p-rh">
                    <div class="profile-label">Relative Humidity (RHobT)</div>
                    {f'<img src="{img_rh}" alt="RH Profile" onclick="openLightbox(this)">' if img_rh else '<div class="no-img">No Image</div>'}
                  </div>
                  <div class="profile-box legend-box">
                    <div class="legend-title">Experiments:</div>
                    {exp_lines}
                    <div class="legend-subtitle">Line Styles:</div>
                    <div class="legend-row">{svg_solid}<span>Solid: RMSE</span></div>
                    <div class="legend-row">{svg_dash}<span>Dash: BIAS</span></div>
                  </div>
                </div>
                '''
            else:
                slide_content = f'''
                <div class="slide-header"><h2>{safe_title}</h2></div>
                <div class="grid-profile-amdar">
                  <div class="profile-box p-temp">
                    <div class="profile-label">Temperature (t)</div>
                    {f'<img src="{img_t}" alt="Temperature Profile" onclick="openLightbox(this)">' if img_t else '<div class="no-img">No Image</div>'}
                  </div>
                  <div class="profile-box p-wind">
                    <div class="profile-label">Wind</div>
                    {f'<img src="{img_w}" alt="Wind Profile" onclick="openLightbox(this)">' if img_w else '<div class="no-img">No Image</div>'}
                  </div>
                  <div class="profile-box legend-box profile-amdar-legend">
                    <div class="legend-title">Experiments:</div>
                    {exp_lines}
                    <div class="legend-subtitle">Line Styles:</div>
                    <div class="legend-row">{svg_solid}<span>Solid: RMSE</span></div>
                    <div class="legend-row">{svg_dash}<span>Dash: BIAS</span></div>
                  </div>
                </div>
                '''

        elif stype == "section_header":
            reg_sub = html.escape(slide.get("subtitle", ""))
            slide_content = f'''
            <div class="section-divider-slide">
              <div class="section-divider-inner">
                <span class="section-pill">REGIONAL VERIFICATION CHAPTER</span>
                <h1 class="section-hero-title">{safe_title}</h1>
                <p class="section-hero-sub">{reg_sub}</p>
              </div>
            </div>
            '''

        elif stype == "radar_case":
            case_t = html.escape(slide.get("case_title", ""))
            cards = slide.get("cards", [])
            
            card_boxes_html = []
            for name, img_path in cards:
                rel_p = rel_img(img_path)
                safe_name = html.escape(name)
                card_boxes_html.append(f'''
                <div class="case-box">
                  <div class="case-label">{safe_name}</div>
                  {f'<img src="{rel_p}" onclick="openLightbox(this)">' if rel_p else '<div class="no-img">No Image</div>'}
                </div>
                ''')

            slide_content = f'''
            <div class="slide-header"><h2>Composite Reflectivity (dBZ) Comparison: {case_t}</h2></div>
            <div class="radar-case-grid">
              {"".join(card_boxes_html)}
            </div>
            '''
            
        elif stype == "summary":
            summary_blocks = []
            for b_text, is_bold in slide.get("bullets", []):
                s_text = html.escape(b_text)
                if is_bold:
                    summary_blocks.append(f'<div class="summary-section-title">{s_text}</div>')
                else:
                    summary_blocks.append(f'<div class="summary-section-body">{s_text}</div>')
            slide_content = f'''
            <div class="slide-header"><h2>{safe_title}</h2></div>
            <div class="summary-card-body">
              {"".join(summary_blocks)}
            </div>
            '''
            
        s_reg = "cases" if slide.get("type") == "radar_case" else slide.get("region", "all")
        slides_html.append(f'''
        <section class="slide-card {'active' if idx == 1 else ''}" id="slide-{idx}" data-slide-num="{idx}" data-region="{s_reg}">
          <div class="slide-watermark">Slide {idx} / {len(slides_data)}</div>
          {slide_content}
        </section>
        ''')

    regions = config.get("regions", ["all"])
    tabs_html = []
    if len(regions) > 1 or any(s.get("type") == "radar_case" for s in slides_data):
        tabs_html.append('<div class="region-tabs" id="regionTabs">')
        tabs_html.append('<button class="region-tab active" data-filter="all" onclick="filterRegion(\'all\', this)">🌐 Full Domain</button>')
        if "east" in regions:
            tabs_html.append('<button class="region-tab" data-filter="east" onclick="filterRegion(\'east\', this)">🗺️ Eastern Region</button>')
        if "west" in regions:
            tabs_html.append('<button class="region-tab" data-filter="west" onclick="filterRegion(\'west\', this)">🏔️ Western Region</button>')
        if any(s.get("type") == "radar_case" for s in slides_data):
            tabs_html.append('<button class="region-tab" data-filter="cases" onclick="filterRegion(\'cases\', this)">📡 Radar Cases</button>')
        tabs_html.append('</div>')
    region_tabs_str = "".join(tabs_html)

    html_template = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>__TITLE__</title>
  <style>
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      background: #0f172a;
      color: #0f172a;
      line-height: 1.4;
      overflow-x: hidden;
    }
    
    /* Top Toolbar */
    .toolbar {
      position: fixed;
      top: 0;
      left: 0;
      right: 0;
      height: 56px;
      background: rgba(16, 44, 87, 0.95);
      backdrop-filter: blur(8px);
      color: #ffffff;
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 0 16px;
      gap: 12px;
      z-index: 1000;
      box-shadow: 0 2px 10px rgba(0,0,0,0.3);
      overflow-x: auto;
    }
    .toolbar::-webkit-scrollbar { display: none; }
    .toolbar-left {
      display: flex;
      align-items: center;
      gap: 10px;
      min-width: 0;
      flex-shrink: 1;
    }
    .toolbar-title {
      font-size: 14.5px;
      font-weight: 700;
      color: #bae6fd;
      white-space: nowrap;
      overflow: hidden;
      text-overflow: ellipsis;
      flex-shrink: 1;
    }
    
    /* Region Navigation Tabs */
    .region-tabs {
      display: flex;
      align-items: center;
      gap: 4px;
      background: rgba(15, 23, 42, 0.6);
      padding: 3px 6px;
      border-radius: 8px;
      border: 1px solid #334155;
      flex-shrink: 0;
    }
    .region-tab {
      background: transparent;
      border: none;
      color: #94a3b8;
      height: 26px;
      padding: 0 10px;
      border-radius: 5px;
      font-size: 12px;
      font-weight: 600;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      transition: all 0.15s ease;
      white-space: nowrap;
      user-select: none;
    }
    .region-tab:hover {
      color: #f8fafc;
      background: rgba(255, 255, 255, 0.08);
    }
    .region-tab.active {
      background: #0284c7;
      color: #ffffff;
      box-shadow: 0 1px 3px rgba(0,0,0,0.2);
    }
    .slide-select {
      background: #1e293b;
      color: #f8fafc;
      border: 1px solid #334155;
      height: 34px;
      padding: 0 10px;
      border-radius: 6px;
      font-size: 12.5px;
      font-weight: 500;
      cursor: pointer;
      outline: none;
      max-width: 260px;
      text-overflow: ellipsis;
      flex-shrink: 1;
      box-sizing: border-box;
    }
    .slide-select:focus { border-color: #0284c7; }
    .toolbar-center {
      display: flex;
      align-items: center;
      gap: 8px;
      font-weight: 600;
      font-size: 13px;
      flex-shrink: 0;
    }
    .counter-badge {
      background: #0284c7;
      color: #ffffff;
      height: 34px;
      padding: 0 12px;
      border-radius: 6px;
      font-size: 12.5px;
      font-weight: 600;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      white-space: nowrap;
      flex-shrink: 0;
      box-sizing: border-box;
    }
    .toolbar-right {
      display: flex;
      align-items: center;
      gap: 8px;
      flex-shrink: 0;
    }
    .tool-btn {
      background: #1e293b;
      color: #e2e8f0;
      border: 1px solid #475569;
      height: 34px;
      padding: 0 12px;
      border-radius: 6px;
      font-size: 12.5px;
      font-weight: 600;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      gap: 6px;
      white-space: nowrap;
      flex-shrink: 0;
      line-height: 1;
      box-sizing: border-box;
      transition: all 0.15s ease;
      user-select: none;
    }
    .tool-btn:hover { background: #334155; color: #ffffff; border-color: #64748b; }
    .tool-btn:focus-visible {
      outline: 2px solid #38bdf8;
      outline-offset: 1px;
    }
    .tool-btn.active {
      background: #0284c7;
      border-color: #38bdf8;
      color: #ffffff;
      box-shadow: 0 1px 3px rgba(0,0,0,0.25);
    }
    .tool-btn.primary { background: #16a34a; border-color: #15803d; color: #fff; }
    .tool-btn.primary:hover { background: #15803d; border-color: #166534; }

    /* Main Container */
    .presentation-container {
      margin-top: 56px;
      padding: 24px 16px;
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 28px;
    }

    /* 16:9 Slide Card */
    .slide-card {
      width: 100%;
      max-width: 1360px;
      aspect-ratio: 16 / 9;
      background: #ffffff;
      border-radius: 10px;
      box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.3), 0 8px 10px -6px rgba(0, 0, 0, 0.2);
      padding: 16px 22px 20px 22px;
      display: flex;
      flex-direction: column;
      position: relative;
      overflow: hidden;
      border: 1px solid #cbd5e1;
      scroll-margin-top: 68px;
    }
    
    .slide-watermark {
      position: absolute;
      bottom: 8px;
      right: 14px;
      font-size: 11px;
      font-weight: 700;
      color: #94a3b8;
      letter-spacing: 0.5px;
      pointer-events: none;
    }

    /* Slideshow mode styles */
    body.mode-slideshow .presentation-container {
      height: calc(100vh - 56px);
      padding: 12px;
      justify-content: center;
      overflow: hidden;
    }
    body.mode-slideshow .slide-card {
      display: none;
      height: calc(100vh - 84px);
      max-height: 85vh;
      width: auto;
    }
    body.mode-slideshow .slide-card.active {
      display: flex;
      animation: fadeIn 0.2s ease-in-out;
    }
    @keyframes fadeIn { from { opacity: 0; transform: scale(0.99); } to { opacity: 1; transform: scale(1); } }

    /* Regional Section Divider Slide */
    .section-divider-slide {
      height: 100%;
      background: linear-gradient(135deg, #102c57 0%, #0d1b3e 100%);
      border-radius: 8px;
      display: flex;
      align-items: center;
      justify-content: center;
      padding: 40px;
      color: #ffffff;
      text-align: center;
      box-shadow: inset 0 0 40px rgba(0,0,0,0.4);
    }
    .section-divider-inner {
      max-width: 800px;
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 16px;
    }
    .section-pill {
      background: rgba(56, 189, 248, 0.15);
      color: #38bdf8;
      border: 1px solid rgba(56, 189, 248, 0.4);
      padding: 5px 14px;
      border-radius: 20px;
      font-size: 12px;
      font-weight: 700;
      letter-spacing: 1.5px;
      text-transform: uppercase;
    }
    .section-hero-title {
      font-size: 32px;
      font-weight: 800;
      color: #ffffff;
      line-height: 1.25;
      text-shadow: 0 2px 10px rgba(0,0,0,0.5);
    }
    .section-hero-sub {
      font-size: 15px;
      color: #cbd5e1;
      line-height: 1.5;
    }

    /* Slide Header */
    .slide-header { text-align: center; margin-bottom: 6px; flex-shrink: 0; }
    .slide-header h2 { font-size: 20px; font-weight: 800; color: #102c57; letter-spacing: -0.3px; }

    /* Title Slide */
    .title-slide-container {
      height: 100%;
      display: flex;
      flex-direction: column;
      justify-content: center;
      padding: 24px 32px;
    }
    .title-hero { margin-bottom: 24px; }
    .report-tag { font-size: 13px; font-weight: 700; color: #0284c7; text-transform: uppercase; letter-spacing: 1.5px; }
    .main-title { font-size: 30px; font-weight: 800; color: #102c57; margin-top: 6px; line-height: 1.25; }
    .title-body { display: flex; gap: 32px; margin-top: 12px; }
    .models-panel { flex: 1.2; background: #f8fafc; border: 1.5px solid #e2e8f0; border-radius: 8px; padding: 18px; }
    .sources-panel { flex: 1.8; background: #f8fafc; border: 1.5px solid #e2e8f0; border-radius: 8px; padding: 18px; }
    .models-panel h3, .sources-panel h3 { font-size: 14.5px; font-weight: 700; color: #1e3a8a; margin-bottom: 12px; }
    .models-grid { display: flex; flex-direction: column; gap: 8px; margin-bottom: 16px; }
    .model-badge { background: #ffffff; padding: 6px 12px; border-radius: 6px; font-size: 12.5px; box-shadow: 0 1px 2px rgba(0,0,0,0.05); }
    .date-ranges { font-size: 12px; color: #475569; display: flex; flex-direction: column; gap: 4px; border-top: 1px dashed #cbd5e1; padding-top: 10px; }
    .bullet-list { list-style: none; display: flex; flex-direction: column; gap: 8px; font-size: 13px; }
    .bullet-header { font-weight: 700; color: #1e293b; margin-top: 4px; }
    .bullet-item { color: #334155; padding-left: 14px; position: relative; }
    .bullet-item::before { content: "•"; position: absolute; left: 2px; color: #0284c7; font-weight: bold; }

    /* 3-Panel Layout */
    .grid-3panel {
      display: grid;
      grid-template-columns: 1fr 1fr;
      grid-template-rows: 1fr 1fr;
      gap: 10px;
      flex: 1;
      min-height: 0;
    }
    .panel-box {
      position: relative;
      background: #fdfdfd;
      border: 1px solid #e2e8f0;
      border-radius: 8px;
      display: flex;
      align-items: center;
      justify-content: center;
      overflow: hidden;
    }
    .panel-box img {
      max-width: 100%;
      max-height: 100%;
      object-fit: contain;
      cursor: zoom-in;
      transition: transform 0.15s ease;
    }
    .panel-box img:hover { transform: scale(1.01); }
    .panel-label {
      position: absolute;
      top: 6px;
      left: 10px;
      background: rgba(16, 44, 87, 0.85);
      color: #ffffff;
      font-size: 11.5px;
      font-weight: 700;
      padding: 3px 8px;
      border-radius: 4px;
      z-index: 2;
    }
    .legend-box {
      padding: 14px 18px;
      display: flex;
      flex-direction: column;
      justify-content: center;
      gap: 6px;
      background: #f8fafc;
      font-size: 12.5px;
      border: 1.5px dashed #cbd5e1;
    }
    .legend-title { font-weight: 700; color: #1e3a8a; margin-bottom: 2px; font-size: 13px; }
    .legend-subtitle { font-weight: 700; color: #1e3a8a; margin-top: 6px; margin-bottom: 2px; font-size: 12px; }
    .legend-row { display: flex; align-items: center; gap: 8px; color: #334155; }
    .color-dot { width: 10px; height: 10px; border-radius: 50%; display: inline-block; flex-shrink: 0; }

    /* Integrated Time Series Layout */
    .ts-integrated-container, .sfc24-container {
      display: flex;
      gap: 12px;
      flex: 1;
      height: calc(100% - 38px);
      overflow: hidden;
    }
    .ts-main-grid {
      display: flex;
      flex-direction: column;
      flex: 1;
      gap: 4px;
      height: 100%;
    }
    .ts-header-row, .sfc24-header-row {
      display: flex;
      text-align: center;
      font-weight: 700;
      font-size: 13px;
      color: #1e3a8a;
      flex-shrink: 0;
    }
    .ts-label-space, .sfc24-label-space { width: 85px; }
    .ts-col-header, .sfc24-col-header { flex: 1; }
    .ts-row, .sfc24-row {
      display: flex;
      align-items: center;
      flex: 1;
      gap: 8px;
      min-height: 0;
    }
    .ts-row-label, .sfc24-label {
      width: 85px;
      font-weight: 700;
      font-size: 12px;
      color: #334155;
      text-align: right;
      padding-right: 6px;
      flex-shrink: 0;
    }
    .ts-img-box, .sfc24-img-box {
      flex: 1;
      height: 100%;
      display: flex;
      align-items: center;
      justify-content: center;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      overflow: hidden;
      background: #f8fafc;
    }
    .ts-img-box img, .sfc24-img-box img {
      max-width: 100%;
      max-height: 100%;
      object-fit: contain;
      cursor: zoom-in;
    }
    .ts-legend-sidebar {
      width: 145px;
      background: #f8fafc;
      border: 1.5px dashed #cbd5e1;
      border-radius: 6px;
      padding: 10px 12px;
      font-size: 11.5px;
      display: flex;
      flex-direction: column;
      gap: 5px;
      flex-shrink: 0;
      height: 100%;
      justify-content: center;
    }

    /* Profile Layout */
    .grid-profile {
      display: grid;
      grid-template-columns: 1fr 1fr;
      grid-template-rows: 1fr 1fr;
      gap: 10px;
      flex: 1;
      min-height: 0;
      height: calc(100% - 38px);
      overflow: hidden;
    }
    .grid-profile-amdar {
      display: flex;
      gap: 14px;
      flex: 1;
      min-height: 0;
      height: calc(100% - 38px);
      overflow: hidden;
    }
    .profile-box {
      position: relative;
      background: #fdfdfd;
      border: 1px solid #e2e8f0;
      border-radius: 8px;
      display: flex;
      align-items: center;
      justify-content: center;
      overflow: hidden;
      min-height: 0;
      min-width: 0;
    }
    .grid-profile-amdar .profile-box {
      flex: 1;
      height: 100%;
    }
    .grid-profile-amdar .profile-amdar-legend {
      flex: 0 0 180px;
      height: 100%;
    }
    .profile-box img { max-width: 100%; max-height: 100%; object-fit: contain; cursor: zoom-in; }
    .profile-label {
      position: absolute;
      top: 6px;
      left: 10px;
      background: rgba(16, 44, 87, 0.85);
      color: #ffffff;
      font-size: 11.5px;
      font-weight: 700;
      padding: 3px 8px;
      border-radius: 4px;
      z-index: 2;
    }

    /* Radar Cases Layout */
    .radar-case-grid {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      grid-template-rows: 1fr 1fr;
      gap: 10px;
      flex: 1;
    }
    .case-box {
      position: relative;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      overflow: hidden;
      display: flex;
      align-items: center;
      justify-content: center;
      background: #f8fafc;
    }
    .case-box img { max-width: 100%; max-height: 100%; object-fit: contain; cursor: zoom-in; }
    .case-label {
      position: absolute;
      top: 4px;
      left: 6px;
      background: rgba(15, 23, 42, 0.85);
      color: #ffffff;
      font-size: 11px;
      font-weight: 700;
      padding: 2px 6px;
      border-radius: 4px;
      z-index: 2;
    }

    /* Summary Slide */
    .summary-card-body {
      padding: 20px 32px;
      display: flex;
      flex-direction: column;
      gap: 14px;
      justify-content: center;
      height: 100%;
    }}
    .summary-section-title { font-size: 15px; font-weight: 700; color: #1e3a8a; border-left: 4px solid #2563eb; padding-left: 8px; }
    .summary-section-body { font-size: 13.5px; color: #334155; padding-left: 14px; line-height: 1.5; }

    /* Lightbox Modal */
    .lightbox-modal {
      display: none;
      position: fixed;
      top: 0; left: 0; right: 0; bottom: 0;
      background: rgba(15, 23, 42, 0.95);
      z-index: 2000;
      align-items: center;
      justify-content: center;
      padding: 20px;
    }
    .lightbox-modal.active { display: flex; }
    .lightbox-modal img {
      max-width: 95vw;
      max-height: 95vh;
      object-fit: contain;
      border-radius: 6px;
      box-shadow: 0 0 30px rgba(0,0,0,0.5);
    }
    .lightbox-close {
      position: absolute;
      top: 20px; right: 24px;
      color: #ffffff;
      font-size: 32px;
      cursor: pointer;
      user-select: none;
    }

    @media print {
      .toolbar { display: none; }
      .presentation-container { margin-top: 0; padding: 0; gap: 0; }
      .slide-card { page-break-after: always; box-shadow: none; border: none; }
    }
  </style>
</head>
<body class="mode-slideshow">

  <!-- Top Navigation Bar -->
  <header class="toolbar">
    <div class="toolbar-left">
      <span class="toolbar-title">⚡ NOAA Verification Report</span>
      __REGION_TABS__
      <select id="slideSelect" class="slide-select" onchange="goToSlide(parseInt(this.value))">
        __NAV_OPTIONS__
      </select>
    </div>
    
    <div class="toolbar-center">
      <button class="tool-btn" onclick="prevSlide()">◀ Prev Slide</button>
      <span class="counter-badge" id="slideCounter">Slide 1 / __TOTAL_SLIDES__</span>
      <button class="tool-btn" onclick="nextSlide()">Next Slide ▶</button>
    </div>
    
    <div class="toolbar-right">
      <button class="tool-btn" id="btnToggleMode" onclick="toggleViewMode()">📜 Continuous Scroll</button>
      <button class="tool-btn" onclick="toggleFullscreen()">⛶ Fullscreen</button>
      <button class="tool-btn primary" onclick="window.print()">🖨️ Export PDF / Print</button>
    </div>
  </header>

  <!-- Presentation Slides -->
  <main class="presentation-container">
    __SLIDES_HTML__
  </main>

  <!-- Image Lightbox Modal -->
  <div id="lightbox" class="lightbox-modal" onclick="closeLightbox()">
    <span class="lightbox-close">&times;</span>
    <img id="lightboxImg" src="" alt="High resolution plot">
  </div>

  <script>
    let currentSlide = 1;
    const totalSlides = __TOTAL_SLIDES__;
    let isSlideshow = true;
    let isProgrammaticScroll = false;

    function updateSlideUI() {
      if (isSlideshow) {
        document.querySelectorAll('.slide-card').forEach(s => s.classList.remove('active'));
        const el = document.getElementById('slide-' + currentSlide);
        if (el) el.classList.add('active');
      } else {
        const el = document.getElementById('slide-' + currentSlide);
        if (el) el.scrollIntoView({ behavior: 'smooth', block: 'start' });
      }
      document.getElementById('slideCounter').innerText = 'Slide ' + currentSlide + ' / ' + totalSlides;
      document.getElementById('slideSelect').value = currentSlide;

      // Sync active region tab with current slide
      const activeSlide = document.getElementById('slide-' + currentSlide);
      if (activeSlide) {
        const currentReg = activeSlide.getAttribute('data-region') || 'all';
        document.querySelectorAll('.region-tab').forEach(b => {
          b.classList.toggle('active', b.getAttribute('data-filter') === currentReg);
        });
      }
    }

    function goToSlide(n) {
      if (n >= 1 && n <= totalSlides) {
        currentSlide = n;
        isProgrammaticScroll = true;
        updateSlideUI();
        setTimeout(() => { isProgrammaticScroll = false; }, 700);
      }
    }

    function filterRegion(reg, btn) {
      if (btn) {
        document.querySelectorAll('.region-tab').forEach(b => b.classList.remove('active'));
        btn.classList.add('active');
      }
      const target = document.querySelector(`.slide-card[data-region="${reg}"]`);
      if (target) {
        const slideNum = parseInt(target.getAttribute('data-slide-num'));
        goToSlide(slideNum);
      }
    }

    function nextSlide() {
      if (currentSlide < totalSlides) {
        currentSlide++;
        updateSlideUI();
      }
    }

    function prevSlide() {
      if (currentSlide > 1) {
        currentSlide--;
        updateSlideUI();
      }
    }

    function toggleViewMode() {
      isSlideshow = !isSlideshow;
      const btn = document.getElementById('btnToggleMode');
      if (isSlideshow) {
        document.body.className = 'mode-slideshow';
        btn.innerText = '📜 Continuous Scroll';
        btn.classList.remove('active');
      } else {
        document.body.className = 'mode-scroll';
        btn.innerText = '🖥️ Slideshow Mode';
        btn.classList.add('active');
      }
      updateSlideUI();
    }

    function toggleFullscreen() {
      if (!document.fullscreenElement) {
        document.documentElement.requestFullscreen().catch(() => {});
      } else {
        document.exitFullscreen().catch(() => {});
      }
    }

    // Auto-update active tab, counter, and select dropdown on scroll
    const scrollObserver = new IntersectionObserver((entries) => {
      if (isSlideshow || isProgrammaticScroll) return;
      entries.forEach(entry => {
        if (entry.isIntersecting) {
          const sNum = parseInt(entry.target.getAttribute('data-slide-num'));
          if (sNum && sNum !== currentSlide) {
            currentSlide = sNum;
            document.getElementById('slideCounter').innerText = 'Slide ' + currentSlide + ' / ' + totalSlides;
            document.getElementById('slideSelect').value = currentSlide;
            const currentReg = entry.target.getAttribute('data-region') || 'all';
            document.querySelectorAll('.region-tab').forEach(b => {
              b.classList.toggle('active', b.getAttribute('data-filter') === currentReg);
            });
          }
        }
      });
    }, {
      rootMargin: '-15% 0px -65% 0px',
      threshold: 0
    });

    document.querySelectorAll('.slide-card').forEach(card => scrollObserver.observe(card));

    // Keyboard navigation
    document.addEventListener('keydown', (e) => {
      if (document.getElementById('lightbox').classList.contains('active')) {
        if (e.key === 'Escape') closeLightbox();
        return;
      }
      if (e.key === 'ArrowRight' || e.key === 'PageDown' || e.key === ' ') {
        e.preventDefault();
        nextSlide();
      } else if (e.key === 'ArrowLeft' || e.key === 'PageUp' || e.key === 'Backspace') {
        e.preventDefault();
        prevSlide();
      } else if (e.key === 'Home') {
        goToSlide(1);
      } else if (e.key === 'End') {
        goToSlide(totalSlides);
      } else if (e.key === 'f' || e.key === 'F') {
        toggleFullscreen();
      }
    });

    // Lightbox functions
    function openLightbox(img) {
      const modal = document.getElementById('lightbox');
      const lightImg = document.getElementById('lightboxImg');
      lightImg.src = img.src;
      modal.classList.add('active');
    }

    function closeLightbox() {
      document.getElementById('lightbox').classList.remove('active');
    }
  </script>
</body>
</html>
"""
    full_html = (
        html_template.replace("__TITLE__", html.escape(title_text))
        .replace("__TOTAL_SLIDES__", str(len(slides_data)))
        .replace("__NAV_OPTIONS__", "".join(nav_options))
        .replace("__SLIDES_HTML__", "".join(slides_html))
        .replace("__REGION_TABS__", region_tabs_str)
    )
    atomic_write_text(output_html, full_html)
    print(f"[DONE] Successfully created HTML presentation: {output_html}")


async def main_async(args):
    config = {}
    if args.config and os.path.exists(args.config):
        with open(args.config, "r") as f:
            config = yaml.safe_load(f)
            
    # Override with CLI args if provided
    if getattr(args, "title", None):
        config["title"] = args.title
    elif "title" not in config:
        config["title"] = "Numerical Weather Prediction Verification Report"
        
    out_dir = args.output_dir or config.get("output_dir", "./output")
    mats_dir = os.path.join(out_dir, "mats_plots")
    os.makedirs(mats_dir, exist_ok=True)
    
    report_name = getattr(args, "report_name", None) or config.get("report_name", "mats_report")
    for ext in [".pptx", ".html", ".ppt"]:
        if report_name.endswith(ext):
            report_name = report_name[:-len(ext)]
    output_pptx = args.output_pptx or config.get("output_pptx", os.path.join(out_dir, f"{report_name}.pptx"))
    output_html = args.output_html or config.get("output_html", os.path.join(out_dir, f"{report_name}.html"))
    
    cref_dir = args.cref_dir or config.get("cref_dir") or os.path.join(out_dir, "cref")
    config["cref_dir"] = cref_dir
    
    extracted_ref_dir = args.extracted_ref_dir or config.get("extracted_ref_dir") or os.path.join(out_dir, "extracted_retv4")
    config["extracted_ref_dir"] = extracted_ref_dir

    sfc_dates = args.surface_dates or config.get("dates", {}).get("surface", "05/06/2024 10:00 - 05/13/2024 00:00")
    ua_dates = args.upperair_dates or config.get("dates", {}).get("upperair", "05/06/2024 00:00 - 05/13/2024 12:00")
    config["dates"] = {"surface": sfc_dates, "upperair": ua_dates}

    # Normalize evaluated regions
    raw_regions = getattr(args, "regions", None) or config.get("regions", ["all"])
    norm_regions = []
    for r in raw_regions:
        r_clean = r.strip().lower()
        if r_clean in ["all", "east", "west"] and r_clean not in norm_regions:
            norm_regions.append(r_clean)
    if "all" not in norm_regions:
        norm_regions.insert(0, "all")
    config["regions"] = norm_regions

    # Process models & color mapping with new flexible parser
    raw_models = []
    if args.models:
        for item in args.models:
            raw_models.append(parse_model_entry(item))
    elif "models" in config:
        for item in config["models"]:
            raw_models.append(parse_model_entry(item))

    if not raw_models:
        print("Error: No models specified. Provide models via --models or config file.")
        sys.exit(1)
        
    models = []
    color_map = {}
    for idx, m in enumerate(raw_models):
        color = m.get("color") or DEFAULT_PALETTE[idx % len(DEFAULT_PALETTE)]
        models.append({
            "id": m["id"],
            "name": m["name"],
            "cluster": m.get("cluster"),
            "path": m.get("path"),
            "color": color
        })
        color_map[m["name"]] = color
        color_map[m["id"]] = color

    config["models"] = models
    config["skip_cref"] = getattr(args, "skip_cref", False)
    force_fetch = getattr(args, "force_fetch", False) or config.get("force_fetch", False)
    config["force_fetch"] = force_fetch

    if args.skip_fetch:
        cache_mode_str = "⚡ SKIP FETCH (Build presentations from local cache only)"
    elif force_fetch:
        cache_mode_str = "🔄 FORCE REFRESH (Override all cache: MATS + Reflectivity)"
    else:
        cache_mode_str = "Standard (Incremental cache with image integrity check)"

    print("=" * 60)
    print(" NOAA GSL Automated Verification & Presentation Tool")
    print("=" * 60)
    print(f" Output Directory : {out_dir}")
    print(f" PPTX Report      : {output_pptx}")
    print(f" HTML Report      : {output_html}")
    print(f" Cache Mode       : {cache_mode_str}")
    print(f" Regions ({len(norm_regions)})   : {', '.join([REGION_NAMES.get(r, r) for r in norm_regions])}")
    print(f" Models ({len(models)}):")
    for m in models:
        path_info = f" [{m['cluster']}:{m['path']}]" if m.get('path') else ""
        print(f"  • {m['name']} -> {m['id']}{path_info}")
    print(f" Surface Dates    : {sfc_dates}")
    print(f" Upper-Air Dates  : {ua_dates}")
    print("=" * 60)

    if getattr(args, "radar_cases", None):
        try:
            if os.path.exists(args.radar_cases):
                with open(args.radar_cases, 'r') as f:
                    config["radar_cases"] = json.load(f)
            else:
                config["radar_cases"] = json.loads(args.radar_cases)
        except Exception as e:
            print(f"  [WARNING] Failed to parse --radar-cases: {e}")

    radar_cases = get_radar_cases(config)

    # Dynamic Stage Planning:
    has_radar = (
        not args.skip_fetch
        and not config.get("skip_cref")
        and bool(radar_cases)
        and any(m.get("path") for m in models)
    )

    if args.skip_fetch:
        stage_defs = [
            ("reports", "Presentation & Reports", 100),
        ]
    elif has_radar:
        stage_defs = [
            ("hpc_radar", "HPC Radar Downloads", 20),
            ("mrms", "MRMS Observations", 20),
            ("surface", "MATS Surface Verification", 20),
            ("upperair", "MATS Upper Air & AMDAR", 15),
            ("refl_vis", "Reflectivity & Cloud/Vis", 13),
            ("reports", "Presentation & Reports", 12),
        ]
    else:
        stage_defs = [
            ("surface", "MATS Surface Verification", 30),
            ("upperair", "MATS Upper Air & AMDAR", 25),
            ("refl_vis", "Reflectivity & Cloud/Vis", 20),
            ("reports", "Presentation & Reports", 25),
        ]

    tracker = PipelineProgressTracker(stage_defs)
    stage_map = {sid: idx for idx, (sid, _, _) in enumerate(stage_defs)}

    # Step 0: Automated synoptic radar image retrieval if paths specified
    if has_radar:
        hpc_idx = stage_map["hpc_radar"]
        fetch_remote_radar_images(
            models, cref_dir, radar_cases=radar_cases, force_fetch=force_fetch,
            on_progress=lambda p: tracker.update(hpc_idx, p)
        )

    # Step 1: Fetch MATS charts and MRMS observations unless --skip-fetch
    if not args.skip_fetch:
        async with async_playwright() as p:
            if has_radar:
                mrms_idx = stage_map["mrms"]
                await fetch_mrms_radar_images(
                    p, cref_dir, radar_cases=radar_cases, force_fetch=force_fetch,
                    on_progress=lambda p: tracker.update(mrms_idx, p)
                )

            sfc_idx = stage_map["surface"]
            tracker.update(sfc_idx, 0)
            await fetch_surface_dieoff(p, models, sfc_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(sfc_idx, 50)
            await fetch_surface_timeseries(p, models, sfc_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(sfc_idx, 100)

            ua_idx = stage_map["upperair"]
            tracker.update(ua_idx, 0)
            await fetch_upperair_dieoff(p, models, ua_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(ua_idx, 35)
            await fetch_upperair_profiles(p, models, ua_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(ua_idx, 70)
            await fetch_upperair_amdar(p, models, ua_dates, color_map, mats_dir, force_fetch=force_fetch)
            tracker.update(ua_idx, 100)

            refl_idx = stage_map["refl_vis"]
            tracker.update(refl_idx, 0)
            await fetch_reflectivity(p, models, sfc_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(refl_idx, 40)
            await fetch_ceiling(p, models, sfc_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(refl_idx, 70)
            await fetch_visibility(p, models, sfc_dates, color_map, mats_dir, regions=config["regions"], force_fetch=force_fetch)
            tracker.update(refl_idx, 100)

    # Step 2: Build PPTX and HTML reports
    rep_idx = stage_map["reports"]
    tracker.update(rep_idx, 5)

    slides_data = []
    if not args.skip_pptx:
        tracker.update(rep_idx, 15)
        slides_data = build_presentation(config, mats_dir, output_pptx)
        tracker.update(rep_idx, 60)
    
    if not args.skip_html:
        if not slides_data:
            slides_data = build_presentation(config, mats_dir, output_pptx)
        tracker.update(rep_idx, 70)
        build_html_report(config, mats_dir, output_html, slides_data)
        tracker.update(rep_idx, 100)

    tracker.complete()



def main():
    parser = argparse.ArgumentParser(description="NOAA GSL Automated Verification Tool & PPT/HTML Generator")
    parser.add_argument("--config", "-c", help="Path to YAML config file")
    parser.add_argument("--title", "-t", help="Title of the presentation report")
    parser.add_argument("--models", "-m", nargs="+", help="Models in format 'ID - cluster:path', 'name:ID', or 'ID'")
    parser.add_argument("--regions", nargs="+", default=["all"], help="Regions to evaluate: all, east, west")
    parser.add_argument("--surface-dates", help="Date range for surface curves (e.g. '05/06/2024 10:00 - 05/13/2024 00:00')")
    parser.add_argument("--upperair-dates", help="Date range for upper air curves")
    parser.add_argument("--output-dir", "-o", help="Directory to save downloaded plots and report")
    parser.add_argument("--report-name", help="Base file name for generated reports (without extension)")
    parser.add_argument("--output-pptx", help="Path of generated PPTX file")
    parser.add_argument("--output-html", help="Path of generated HTML presentation file")
    parser.add_argument("--cref-dir", help="Path to composite reflectivity images directory")
    parser.add_argument("--extracted-ref-dir", help="Path to extracted reference images directory")
    parser.add_argument("--radar-cases", help="JSON string or file path specifying radar evaluation cases")
    parser.add_argument("--skip-cref", action="store_true", help="Skip radar composite reflectivity comparisons")
    parser.add_argument("--skip-fetch", action="store_true", help="Skip fetching images (use existing images to build PPT & HTML)")
    parser.add_argument("--force-fetch", action="store_true", help="Force re-fetching images from MATS and HPC, bypassing local cache")
    parser.add_argument("--skip-pptx", action="store_true", help="Skip generating PPTX")
    parser.add_argument("--skip-html", action="store_true", help="Skip generating HTML report")
    
    args = parser.parse_args()
    asyncio.run(main_async(args))

if __name__ == "__main__":
    main()
